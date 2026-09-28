"""Write the Category III rows as CSVs and one workbook.

=====================  =========================================================
Sheet                  Contents
=====================  =========================================================
``README``             provenance, column notes, the quirks of the source page
``CatIII_NetFigures``  the headline table, every quarter since September 2012
``CatIII_EquityDebt``  the equity and debt table, every quarter since March 2024
``CatIII_Combined``    both side by side, one row per quarter
=====================  =========================================================

Rows run oldest first so the sheets chart and pivot directly.
"""
from __future__ import annotations

import calendar
import csv
import logging
import os
import re
from datetime import date, datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .parse import EQUITY_DEBT, NET_FIGURES, URL, Quarter

log = logging.getLogger(__name__)

WORKBOOK = "sebi_aif_category3.xlsx"
UNIT = " (Rs crore)"
ID_COLS = ["Quarter End", "Quarter", "Financial Year", "FY Quarter"]

_HDR_FILL = PatternFill("solid", fgColor="1F3864")
_HDR_FONT = Font(color="FFFFFF", bold=True, size=10)


def financial_year(d: date) -> str:
    """India's April-March financial year a quarter end falls in, e.g. '2026-27'."""
    start = d.year if d.month >= 4 else d.year - 1
    return f"{start}-{(start + 1) % 100:02d}"


def fy_quarter(d: date) -> str:
    """Quarter of the financial year: June is Q1, March is Q4."""
    return f"Q{(d.month - 4) % 12 // 3 + 1}"


def _ident(q: Quarter) -> list:
    d = q.quarter_end
    return [d, f"{calendar.month_abbr[d.month]}-{d.year}", financial_year(d), fy_quarter(d)]


def _notes(q: Quarter) -> str:
    return "; ".join(q.notes)


def _columns(dicts, spec) -> list[str]:
    """Known columns in SEBI's order, then any column SEBI adds later."""
    cols = list(spec)
    for d in dicts:
        cols += [k for k in d if k not in cols]
    return cols


def _titles(cols: list[str]) -> list[str]:
    """Known columns carry the unit; one SEBI adds later keeps its printed header,
    since nothing says it is an amount at all."""
    known = {*NET_FIGURES, *EQUITY_DEBT}
    return [c + UNIT if c in known else c for c in cols]


def build_tables(quarters: list[Quarter]) -> dict[str, tuple[list[str], list[list]]]:
    """Header and rows for each output table."""
    net_cols = _columns([q.net for q in quarters], NET_FIGURES)
    eq_quarters = [q for q in quarters if q.equity_debt is not None]
    eq_cols = _columns([q.equity_debt for q in eq_quarters], EQUITY_DEBT)
    trail = ["SEBI Row Label", "SEBI Section Heading", "Notes"]

    net = ([*ID_COLS, *_titles(net_cols), *trail],
           [[*_ident(q), *(q.net.get(c) for c in net_cols), q.net_label, q.heading, _notes(q)]
            for q in quarters])
    eq = ([*ID_COLS, *_titles(eq_cols), *trail],
          [[*_ident(q), *(q.equity_debt.get(c) for c in eq_cols), q.equity_debt_label,
            q.heading, _notes(q)]
           for q in eq_quarters])
    combined = ([*ID_COLS, *_titles(net_cols + eq_cols), *trail[1:]],
                [[*_ident(q), *(q.net.get(c) for c in net_cols),
                  *((q.equity_debt or {}).get(c) for c in eq_cols), q.heading, _notes(q)]
                 for q in quarters])
    return {"CatIII_NetFigures": net, "CatIII_EquityDebt": eq, "CatIII_Combined": combined}


def _span(qs: list[Quarter]) -> str:
    return f"{qs[0].quarter_end:%B %Y} to {qs[-1].quarter_end:%B %Y}"


def _readme(quarters: list[Quarter], tables, generated: str) -> list[tuple[str, bool]]:
    eq = [q for q in quarters if q.equity_debt is not None]
    redated = [q for q in quarters if q.stated_date != q.quarter_end]
    footnote = next((q.footnote for q in reversed(quarters) if q.footnote), None)

    lines = [
        ("SEBI - Data relating to activities of Alternative Investment Funds: Category III AIFs", True),
        ("", False),
        ("Source", True),
        (URL, False),
        (f"Generated (UTC): {generated}", False),
        ("", False),
        ("Sheets", True),
        ("CatIII_NetFigures   'Cumulative net figures as at the end of the quarter ending ...'"
         " - the Category III row, every column", False),
        ("CatIII_EquityDebt   'Cumulative net investment made in equity and debt securities'"
         " - the Category III row, every column", False),
        ("CatIII_Combined     both tables side by side, one row per quarter", False),
        ("", False),
        ("Coverage", True),
        (f"Net figures table: {len(quarters)} quarters, {_span(quarters)} - every quarter on the page.",
         False),
    ]
    if eq:
        lines.append((f"Equity and debt table: {len(eq)} quarters, {_span(eq)} - SEBI publishes it"
                      f" only from {eq[0].quarter_end:%B %Y}, so it is blank in CatIII_Combined"
                      " before then.", False))
    lines += [
        ("", False),
        ("Notes", True),
        ("Figures are cumulative net amounts in Rs crore as at the quarter end, exactly as SEBI"
         " publishes them.", False),
        ("Precision follows SEBI: whole crores from March 2024, up to three decimals before.", False),
        ("A blank figure means SEBI printed '-' (nil) or left the cell empty.", False),
        ("Before March 2024 the heading reads 'Cumulative net figures as at the end of <date>'"
         " - same table, same columns.", False),
        ("Financial Year / FY Quarter use India's April-March year (quarter ending June = Q1).",
         False),
    ]
    if redated:
        lines.append((f"{len(redated)} SEBI headings repeat the quarter above them; those sections"
                      " are dated by position and flagged in the Notes column:", False))
        lines += [(f"    '{q.heading}'  ->  {q.quarter_end:%B %d, %Y}", False) for q in redated]
    if footnote:
        lines += [("", False), ("SEBI's definitions for the equity and debt table", True)]
        lines += [(part.strip(), False) for part in re.split(r"\s(?=#)", footnote)]
    lines += [("", False), ("Row counts", True)]
    lines += [(f"{name}: {len(rows)} rows", False) for name, (_, rows) in tables.items()]
    return lines


def _csv_value(v):
    return int(v) if isinstance(v, float) and v.is_integer() else v


def write_outputs(quarters: list[Quarter], out_dir: str) -> dict[str, int]:
    """Write ``<out_dir>/csv/*.csv`` and ``<out_dir>/sebi_aif_category3.xlsx``."""
    tables = build_tables(quarters)
    generated = datetime.now(timezone.utc).isoformat(timespec="seconds")

    csv_dir = os.path.join(out_dir, "csv")
    os.makedirs(csv_dir, exist_ok=True)
    for name, (head, rows) in tables.items():
        with open(os.path.join(csv_dir, f"{name}.csv"), "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(head)
            w.writerows([_csv_value(v) for v in r] for r in rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for text, bold in _readme(quarters, tables, generated):
        ws.append([text])
        if bold:
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)
    ws.column_dimensions["A"].width = 120

    for name, (head, rows) in tables.items():
        ws = wb.create_sheet(name)
        ws.append(head)
        for r in rows:
            ws.append(r)
        for idx, title in enumerate(head, 1):
            letter = get_column_letter(idx)
            cell = ws.cell(1, idx)
            cell.fill, cell.font = _HDR_FILL, _HDR_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            fmt, width = None, 14
            if title == "Quarter End":
                fmt, width = "yyyy-mm-dd", 12
            elif title.endswith(UNIT):
                fmt, width = "#,##0.00", 18
            elif title == "SEBI Section Heading":
                width = 60
            elif title == "Notes":
                width = 50
            ws.column_dimensions[letter].width = width
            if fmt:
                for (c,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    c.number_format = fmt
        ws.row_dimensions[1].height = 45
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

    path = os.path.join(out_dir, WORKBOOK)
    wb.save(path)
    log.info("wrote %s and %d CSVs in %s", path, len(tables), csv_dir)
    return {name: len(rows) for name, (_, rows) in tables.items()}
