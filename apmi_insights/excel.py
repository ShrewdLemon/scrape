"""Turn action rows into the table as the site displays it, then into Excel.

Sheet 1 mirrors the on-screen columns in order.  The site's first column,
"IA Details", stacks the approach name over the portfolio manager's name; it is
split into two sub-columns here so each stays filterable.  "Actions" is the
row's "View Details" link.  The unlabeled expand-arrow column has no data and
is left out.

Numbers stay numbers.  Returns keep the API's full precision; the cell format
shows one decimal and a % sign, as the site does (``90.09`` -> ``90.1%``).
The site turns a missing return into ``0.0%`` (``parseFloat("" || 0)``).  Here
a missing value is left blank so it can't be mistaken for a real 0 %.
"""
from __future__ import annotations

from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .fetch import BASE

_HDR_FILL = PatternFill("solid", fgColor="1F3864")
_HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
_LINK_FONT = Font(color="0563C1", underline="single")

PCT = '0.0"%";-0.0"%";0.0"%"'
AUM = '#,##0.00" Cr"'
DATE = "dd-mm-yyyy"


def columns(period: str = "SI") -> list[tuple[str, int, str | None]]:
    """(header, width, number format) in on-screen order."""
    return [
        ("IA Details - Investment Approach", 48, None),
        ("IA Details - Portfolio Manager", 48, None),
        ("Service Type", 13, None),
        ("AUM", 16, AUM),
        ("Inception Date", 15, DATE),
        (f"IA({period})", 11, PCT),
        (f"BENCHMARK({period})", 16, PCT),
        ("Actions", 14, None),
    ]


PERIOD_FIELDS = {"SI": ("iaSinceInception", "bchSinceInception")}


def number(v) -> float | None:
    if v is None:
        return None
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "NA", "N/A", "null"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def inception(v) -> date | None:
    try:
        return datetime.strptime(str(v).strip(), "%d-%m-%Y").date()
    except ValueError:
        return None


def detail_url(row: dict) -> str:
    return f"{BASE}/ia-insight-report/{row['pmsIaId']}"


def to_table(rows: list[dict], period: str = "SI") -> list[list]:
    ia_key, bch_key = PERIOD_FIELDS[period]
    out = []
    for r in rows:
        out.append([
            r.get("iaName"),
            r.get("pmsProviderName"),
            "D" if r.get("iaServiceType") == "D" else "ND",
            number(r.get("aum")),
            inception(r.get("iaDateOfInception")) or (r.get("iaDateOfInception") or None),
            number(r.get(ia_key)),
            number(r.get(bch_key)),
            detail_url(r),
        ])
    return out


def write_workbook(path: str, rows: list[dict], meta: list[tuple[str, object]],
                   period: str = "SI", title: str = "Equity SI") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    cols = columns(period)

    for c, (head, width, _) in enumerate(cols, 1):
        cell = ws.cell(1, c, head)
        cell.fill, cell.font = _HDR_FILL, _HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 30

    link_col = len(cols)
    for r, values in enumerate(to_table(rows, period), 2):
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            fmt = cols[c - 1][2]
            if fmt and v is not None:
                cell.number_format = fmt
            if c == link_col:
                cell.value, cell.hyperlink, cell.font = "View Details", v, _LINK_FONT
            elif c == 3:
                cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 1)}"

    info = wb.create_sheet("Run Info")
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 90
    for r, (k, v) in enumerate(meta, 1):
        info.cell(r, 1, k).font = Font(bold=True)
        info.cell(r, 2, v).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)
