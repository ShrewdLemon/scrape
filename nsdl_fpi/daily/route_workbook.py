"""Four-sheet workbook: Total, Stock Exchange, Primary, and a formula-driven Combined.

The three data sheets hold NSDL's figures as published; ``Combined`` holds
no numbers of its own, only formulas over them:

* ``Total``          - equity per reporting day, 2001 onwards (the single Equity
                       row before Dec 2009, the Equity Sub-total after).
* ``Stock Exchange`` - the Stock Exchange equity route (Dec 2009 onwards).
* ``Primary``        - the Primary market & others equity route (Dec 2009 onwards).
* ``Combined``       - one row per Total row; links to Total, looks each date up in
                       the two route sheets (MATCH once per row, then INDEX), checks
                       Stock Exchange + Primary against Total, and summarises by year
                       with SUMIFS.

All three data sheets share one column layout (A Date ... J USD/INR rate), so
every Combined formula reads the same column letter from each. ``FY`` is the
Indian financial year (April-March), e.g. 15-Jan-2010 -> "FY 2009-10", computed
by formula from the date.
"""
from __future__ import annotations

import os
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from .fetch import URL
from .parse import MonthReport

FONT = "Arial"
F_TITLE = Font(name=FONT, bold=True, size=13)
F_NOTE = Font(name=FONT, italic=True, size=9, color="595959")
F_HEAD = Font(name=FONT, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, color="0000FF")      # values scraped from NSDL
F_LINK = Font(name=FONT, color="008000")       # links to another sheet
F_CALC = Font(name=FONT, color="000000")       # formulas on this sheet
F_BOLD = Font(name=FONT, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="1F4E78")
FILL_GROUP = {"Total": "2E75B6", "Stock Exchange": "548235", "Primary": "BF8F00",
              "Check": "7F7F7F", "Year": "1F4E78"}
FILL_ASSUMPTION = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")

NUM = '#,##0.00;[Red](#,##0.00);"-"'
RATE = "0.0000"
DATE = "dd-mmm-yyyy"

HEADER_ROW, FIRST = 4, 5           # header row / first data row on every sheet
DATA_HEAD = ["Date", "Year", "Month", "FY", None, "Gross Purchases (Rs Crore)", "Gross Sales (Rs Crore)",
             "Net Investment (Rs Crore)", "Net Investment (US$ Million)", "USD/INR Rate"]
VAL = "FGHI"                       # the four value columns on the data sheets
RATE_COL = "J"


def fy_formula(cell: str) -> str:
    """Indian financial year (Apr-Mar) of a date cell, as text: "FY 2009-10"."""
    return (f'=IF(MONTH({cell})>=4,"FY "&YEAR({cell})&"-"&RIGHT(YEAR({cell})+1,2),'
            f'"FY "&(YEAR({cell})-1)&"-"&RIGHT(YEAR({cell}),2))')


def fy_label(d) -> str:
    y = d.year if d.month >= 4 else d.year - 1
    return f"FY {y}-{str(y + 1)[-2:]}"
ROUTES = {"Stock Exchange": "Stock Exchange", "Primary": "Primary market & others"}


def _head(ws, row, labels, fill=FILL_HEAD, height=32):
    for i, lab in enumerate(labels, 1):
        if lab is None:
            continue
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill = F_HEAD, fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height


def _title(ws, title, note):
    ws["A1"], ws["A2"] = title, note
    ws["A1"].font, ws["A2"].font = F_TITLE, F_NOTE


def _data_sheet(wb, name, title, note, label, rows):
    """rows: (date, label_value, gp, gs, net_inr, net_usd, fx)."""
    ws = wb.create_sheet(name)
    _title(ws, title, note)
    head = list(DATA_HEAD)
    head[4] = label
    _head(ws, HEADER_ROW, head)
    for i, (d, lab, *vals) in enumerate(rows, FIRST):
        ws.cell(row=i, column=1, value=d).number_format = DATE
        ws.cell(row=i, column=2, value=d.year)
        ws.cell(row=i, column=3, value=d.month)
        ws.cell(row=i, column=5, value=lab)
        for j, v in enumerate(vals, 6):
            ws.cell(row=i, column=j, value=v).number_format = RATE if j == 10 else NUM
        for j in range(1, 11):
            ws.cell(row=i, column=j).font = F_INPUT
        fy = ws.cell(row=i, column=4, value=fy_formula(f"A{i}"))
        fy.font, fy.alignment = F_CALC, Alignment(horizontal="center")
    for col, w in zip("ABCDEFGHIJ", (13, 7, 7, 11, 18, 15, 15, 15, 15, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=FIRST, column=2)
    last = FIRST + len(rows) - 1
    ws.auto_filter.ref = f"A{HEADER_ROW}:J{max(last, FIRST)}"
    return last


def write_route_workbook(path: str, reports: list[MonthReport], note: str = "") -> dict:
    """``note`` is appended to the provenance line on every sheet (e.g. merge sources)."""
    reports = sorted(reports, key=lambda r: (r.year, r.month))
    days = [d for r in reports for d in r.days]
    routed = [d for d in days if d.routes]
    split_from = routed[0].day if routed else None
    stamp = (f"Source: NSDL Archive (Daily Trends in FPI Investments), {URL} - scraped "
             f"{datetime.now(timezone.utc):%Y-%m-%d} UTC. "
             f"Coverage {days[0].day:%d %b %Y} to {days[-1].day:%d %b %Y}." + (f" {note}" if note else ""))

    wb = Workbook()
    wb.remove(wb.active)

    # 1. Total ---------------------------------------------------------------
    last_t = _data_sheet(
        wb, "Total", "FPI investment in equity - TOTAL, per reporting day",
        stamp + (f" Before {split_from:%d %b %Y} NSDL published one undivided Equity row per day"
                 " (Basis = Equity); from then on this is the Equity Sub-total"
                 " (= Stock Exchange + Primary market & others)." if split_from else ""),
        "Basis",
        [(d.day, d.basis, d.gross_purchases, d.gross_sales, d.net_inr, d.net_usd, d.fx) for d in days])

    # 2./3. Routes -----------------------------------------------------------
    last_r = {}
    for sheet, route in ROUTES.items():
        last_r[sheet] = _data_sheet(
            wb, sheet, f"FPI investment in equity - {route.upper()} route, per reporting day",
            stamp + f" NSDL splits equity by investment route from {split_from:%d %b %Y}"
            if split_from else stamp,
            "Investment Route",
            [(d.day, route, *d.routes[route], d.fx) for d in routed if route in d.routes])

    # 4. Combined ------------------------------------------------------------
    # A Date | B Year | C Month | D FY | E split? | F:I Total | J:M Stock Exchange |
    # N:Q Primary | R SE+Primary | S diff | T check | U rate | V:W hidden helpers
    ws = wb.create_sheet("Combined")
    _title(ws, "Combined - Total vs Stock Exchange + Primary (all formulas)",
           "Black = formula, green = link to another sheet, yellow = editable tolerance. "
           "FY = Indian financial year (Apr-Mar). Route columns are blank before NSDL started "
           "publishing the split. Hidden helper columns V:W hold each date's row in the two route sheets.")
    ws["R2"], ws["T2"] = "Check tolerance (Rs Cr):", 0.1
    ws["R2"].font = F_BOLD
    ws["R2"].alignment = Alignment(horizontal="right")
    ws["T2"].font, ws["T2"].fill = F_INPUT, FILL_ASSUMPTION
    ws["T2"].comment = Comment("Largest |Total - (Stock Exchange + Primary)| still treated as OK. "
                               "NSDL rounds each figure to 2 decimals, so gaps of 0.01-0.10 are rounding.",
                               "scraper")
    tol = "$T$2"

    def band(label, c0, c1, key):
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c1)
        c = ws.cell(row=3, column=c0, value=label)
        c.font, c.fill = F_HEAD, PatternFill("solid", fgColor=FILL_GROUP[key])
        c.alignment = Alignment(horizontal="center")

    band("TOTAL", 6, 9, "Total")
    band("STOCK EXCHANGE", 10, 13, "Stock Exchange")
    band("PRIMARY MARKET & OTHERS", 14, 17, "Primary")
    band("CHECK", 18, 21, "Check")
    measures = ["Gross Purchases (Rs Cr)", "Gross Sales (Rs Cr)", "Net (Rs Cr)", "Net (US$ mn)"]
    head = (["Date", "Year", "Month", "FY", "Route split published?"] + measures * 3 +
            ["Stock Exch + Primary Net (Rs Cr)", "Total Net - (SE + Primary) (Rs Cr)", "Check",
             "USD/INR Rate", "SE row (helper)", "Primary row (helper)"])
    _head(ws, HEADER_ROW, head, height=45)

    TOT, SE, PM = "FGHI", "JKLM", "NOPQ"
    se_last, pm_last = last_r["Stock Exchange"], last_r["Primary"]
    se_rng = f"'Stock Exchange'!$A${FIRST}:$A${se_last}"
    pm_rng = f"Primary!$A${FIRST}:$A${pm_last}"
    links = set("A" + TOT + SE + PM + "U")
    for r in range(FIRST, last_t + 1):
        f = {
            "A": f"=Total!A{r}",
            "B": f"=YEAR(A{r})",
            "C": f"=MONTH(A{r})",
            "D": fy_formula(f"A{r}"),
            "E": f'=IF(AND(V{r}<>"",W{r}<>""),"Yes","No")',
            "R": f'=IF($E{r}="Yes",L{r}+P{r},"")',
            "S": f'=IF($E{r}="Yes",H{r}-R{r},"")',
            "T": f'=IF($E{r}="Yes",IF(ROUND(ABS(S{r}),2)<={tol},"OK","CHECK"),"n/a")',
            "U": f"=Total!{RATE_COL}{r}",
            "V": f'=IFERROR(MATCH($A{r},{se_rng},0),"")',
            "W": f'=IFERROR(MATCH($A{r},{pm_rng},0),"")',
        }
        for k, src in enumerate(VAL):
            f[TOT[k]] = f"=Total!{src}{r}"
            f[SE[k]] = f"=IF($V{r}=\"\",\"\",INDEX('Stock Exchange'!{src}${FIRST}:{src}${se_last},$V{r}))"
            f[PM[k]] = f'=IF($W{r}="","",INDEX(Primary!{src}${FIRST}:{src}${pm_last},$W{r}))'
        for col, formula in f.items():
            c = ws[f"{col}{r}"]
            c.value = formula
            c.font = F_LINK if col in links else F_CALC
            if col == "A":
                c.number_format = DATE
            elif col == "U":
                c.number_format = RATE
            elif col in TOT + SE + PM + "RS":
                c.number_format = NUM
            elif col in "DET":
                c.alignment = Alignment(horizontal="center")

    widths = {"A": 13, "B": 7, "C": 7, "D": 11, "E": 11, "R": 15, "S": 15, "T": 9, "U": 10,
              "V": 9, "W": 9, "X": 3}
    for i in range(1, 25):
        col = get_column_letter(i)
        ws.column_dimensions[col].width = widths.get(col, 13)
    ws.column_dimensions["V"].hidden = ws.column_dimensions["W"].hidden = True
    ws.freeze_panes = ws.cell(row=FIRST, column=6)
    ws.auto_filter.ref = f"A{HEADER_ROW}:U{last_t}"

    # Summaries to the right: by calendar year, then by FY (SUMIFS/COUNTIFS) -
    rng = lambda col: f"${col}${FIRST}:${col}${last_t}"   # noqa: E731

    def summary(c0, title, key_label, key_col, keys):
        sh = [key_label, "Reporting days", "Days with route split", "Total Net (Rs Cr)",
              "Total Net on split days (Rs Cr)", "Stock Exch Net (Rs Cr)", "Primary Net (Rs Cr)",
              "SE + Primary Net (Rs Cr)", "Difference (Rs Cr)", "Total Net (US$ mn)"]
        L = [get_column_letter(c0 + i) for i in range(len(sh))]
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + len(sh) - 1)
        c = ws.cell(row=3, column=c0, value=title)
        c.font, c.fill = F_HEAD, PatternFill("solid", fgColor=FILL_GROUP["Year"])
        c.alignment = Alignment(horizontal="center")
        for i, lab in enumerate(sh):
            c = ws.cell(row=HEADER_ROW, column=c0 + i, value=lab)
            c.font, c.fill = F_HEAD, FILL_HEAD
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for k, key in enumerate(keys):
            r = FIRST + k
            ws[f"{L[0]}{r}"] = key
            ws[f"{L[0]}{r}"].font = F_INPUT
            by = f"{rng(key_col)},{L[0]}{r}"
            cells = [f"=COUNTIFS({by})",
                     f'=COUNTIFS({by},{rng("E")},"Yes")',
                     f"=SUMIFS({rng('H')},{by})",
                     f'=SUMIFS({rng("H")},{by},{rng("E")},"Yes")',
                     f"=SUMIFS({rng('L')},{by})",
                     f"=SUMIFS({rng('P')},{by})",
                     f"={L[5]}{r}+{L[6]}{r}",
                     f"={L[4]}{r}-{L[7]}{r}",
                     f"=SUMIFS({rng('I')},{by})"]
            for i, formula in enumerate(cells, 1):
                c = ws[f"{L[i]}{r}"]
                c.value, c.font = formula, F_CALC
                c.number_format = "#,##0" if i <= 2 else NUM
        tr = FIRST + len(keys)
        ws[f"{L[0]}{tr}"] = "Total"
        for i in range(len(sh)):
            c = ws[f"{L[i]}{tr}"]
            if i:
                c.value = f"=SUM({L[i]}{FIRST}:{L[i]}{tr - 1})"
                c.number_format = "#,##0" if i <= 2 else NUM
            c.font = F_BOLD
            c.border = Border(top=THIN, bottom=Side(style="double", color="000000"))
        note = ws.cell(row=tr + 2, column=c0,
                       value=f"Data covers {days[0].day:%d %b %Y} to {days[-1].day:%d %b %Y}; "
                             "the first and last periods may be partial.")
        note.font = F_NOTE
        for i, col in enumerate(L):
            ws.column_dimensions[col].width = 11 if i == 0 else 14
        return c0 + len(sh)

    years = sorted({d.day.year for d in days})
    fys = sorted({fy_label(d.day) for d in days})
    nxt = summary(25, "SUMMARY BY CALENDAR YEAR", "Year", "B", years)      # from column Y
    ws.column_dimensions[get_column_letter(nxt)].width = 3
    summary(nxt + 1, "SUMMARY BY FINANCIAL YEAR (Apr-Mar)", "FY", "D", fys)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    return {"total_rows": last_t - FIRST + 1,
            "route_rows": {k: v - FIRST + 1 for k, v in last_r.items()},
            "years": len(years), "fys": len(fys)}
