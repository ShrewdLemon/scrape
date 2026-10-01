"""Workbook (and CSV) output for the daily FPI equity series."""
from __future__ import annotations

import csv
import os
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..excel import BAD, FILL, HEAD, NUM, _header
from ..parse import MONTHS
from .fetch import URL
from .parse import EQUITY_BASES, ROUTE_NAMES, VALUE_COLS, MonthReport

DEC = "#,##0.00;[Red]-#,##0.00"
FALLBACK = PatternFill("solid", fgColor="FFF2CC")
TOL = 1.0          # crore / USD mn; NSDL rounds every daily figure to 2 dp

DAILY_COLS = ["Date", "Year", "Month", "Basis", "Gross Purchases (Rs Crore)",
              "Gross Sales (Rs Crore)", "Net Investment (Rs Crore)",
              "Net Investment (US$ Million)", "USD/INR Rate"]
MEASURES = ["Gross Purchases (Rs Crore)", "Gross Sales (Rs Crore)",
            "Net Investment (Rs Crore)", "Net Investment (US$ Million)"]
ROUTE_COLS = ["Date", "Year", "Month"] + [f"{r} - {m}" for r in ROUTE_NAMES for m in MEASURES] + ["USD/INR Rate"]


def daily_rows(reports: list[MonthReport]) -> list[dict]:
    out = []
    for r in sorted(reports, key=lambda r: (r.year, r.month)):
        for d in r.days:
            out.append(dict(zip(DAILY_COLS, [d.day, d.day.year, d.day.month, d.basis,
                                             d.gross_purchases, d.gross_sales, d.net_inr,
                                             d.net_usd, d.fx])))
    return out


def route_rows(reports: list[MonthReport]) -> list[dict]:
    """Wide rows: the three equity routes per day, routed-layout days only."""
    out = []
    for r in sorted(reports, key=lambda r: (r.year, r.month)):
        for d in r.days:
            if not d.routes:
                continue
            vals = [v for name in ROUTE_NAMES for v in d.routes.get(name, (None,) * len(VALUE_COLS))]
            out.append(dict(zip(ROUTE_COLS, [d.day, d.day.year, d.day.month, *vals, d.fx])))
    return out


def check_routes(r: MonthReport) -> tuple[float | None, bool]:
    """(largest |diff|, ok) of each route's daily sum vs NSDL's month total, all four columns."""
    if not any(d.routes for d in r.days):
        return None, True
    if set(r.stated_routes) != set(ROUTE_NAMES):
        return None, False
    worst = max(abs(round(r.route_sum(name, i) - (r.stated_routes[name][i] or 0), 2))
                for name in ROUTE_NAMES for i in range(len(VALUE_COLS)))
    return worst, worst <= TOL


def check(r: MonthReport) -> tuple[float | None, float | None, bool]:
    """(diff INR, diff USD, ok) of the summed daily equity vs NSDL's month total."""
    if not r.stated or r.stated[2] is None:
        return None, None, False
    di = round(r.computed_net_inr - r.stated[2], 2)
    du = None if r.stated[3] is None else round(r.computed_net_usd - r.stated[3], 2)
    return di, du, abs(di) <= TOL and (du is None or abs(du) <= TOL)


def mismatches(reports: list[MonthReport]) -> list[MonthReport]:
    return [r for r in reports
            if (not r.fallback_days and not check(r)[2]) or not check_routes(r)[1]]


def write_csv(path: str, rows: list[dict], cols: list[str] = DAILY_COLS) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({**r, "Date": r["Date"].isoformat()})


def write_workbook(path: str, reports: list[MonthReport]) -> None:
    reports = sorted(reports, key=lambda r: (r.year, r.month))
    rows = daily_rows(reports)
    rroutes = route_rows(reports)
    fallbacks = [r for r in rows if r["Basis"] not in EQUITY_BASES]
    wb = Workbook()

    # README -----------------------------------------------------------------
    ws = wb.active
    ws.title = "README"
    first, last = rows[0]["Date"], rows[-1]["Date"]
    routed = [r for r in reports if r.layout == "routed"]
    lines = [
        ("FPI Investment in Equity - daily", True),
        ("", False),
        (f"Source: {URL}", False),
        ("Report: NSDL 'Archive (Trends in FPI/FII Investments)' - Daily Trends, Equity row only.", False),
        (f"Coverage: {first:%d %b %Y} to {last:%d %b %Y} - {len(rows)} reporting days in {len(reports)} months.", False),
        (f"Scraped: {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC", False),
        ("", False),
        ("Which row is 'Equity'", True),
        ("Until the layout change each day has a single Equity row (Basis = 'Equity').", False),
        (("From " + f"{routed[0].year}-{routed[0].month:02d}" if routed else "Later") +
         " equity is split into Stock Exchange / Primary market & others; the Equity Sub-total is used"
         " (Basis = 'Equity sub-total').", False),
        ("If a day has no equity figures, its Total row is used instead (Basis = 'Total row'),", False),
        ("highlighted yellow and listed on the Fallbacks sheet. "
         f"Days on a fallback basis: {len(fallbacks)}.", False),
        ("", False),
        ("Units: Rs Crore and US$ Million as NSDL publishes them; rate = NSDL's USD/INR conversion for the day.", False),
        ("Net = gross purchases - gross sales. Dates are NSDL 'reporting dates' (trades up to the previous day).", False),
        ("", False),
        ("Sheets", True),
        ("Equity_Daily     one row per reporting day (filterable)", False),
        ("Monthly          daily figures summed per month, next to NSDL's own 'Total for <Month>' equity row", False),
        ("INR_by_Year      net equity (Rs Crore) per month, years down, with an annual chart", False),
        ("Equity_Routes    equity split by investment route per day: Stock Exchange / Primary market & others /", False),
        ("                 Sub-total (each with gross purchases, gross sales, net Rs Cr, net US$ mn)", False),
        ("Routes_Monthly   each route's daily sum vs NSDL's own month total for that route", False),
        ("Fallbacks        days where the Total row stood in for Equity", False),
        ("", False),
        ("Investment routes", True),
        (("NSDL splits equity by route only from " + (f"{routed[0].year}-{routed[0].month:02d}" if routed else "the routed layout") +
          f"; before that each day has one undivided Equity row, so Equity_Routes starts there ({len(rroutes)} days)."), False),
        ("Stock Exchange + Primary market & others = Sub-total (to rounding); the Sub-total is the Equity figure on Equity_Daily.", False),
    ]
    for text, bold in lines:
        ws.append([text])
        if bold:
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=13 if ws.max_row == 1 else 11)
    ws.column_dimensions["A"].width = 110

    # Daily ------------------------------------------------------------------
    ws = wb.create_sheet("Equity_Daily")
    _header(ws, DAILY_COLS, {1: 12, 2: 7, 3: 7, 4: 18, 5: 16, 6: 16, 7: 16, 8: 16, 9: 10})
    for r in rows:
        ws.append([r[c] for c in DAILY_COLS])
        row = ws[ws.max_row]
        row[0].number_format = "dd-mmm-yyyy"
        for c in row[4:8]:
            c.number_format = DEC
        row[8].number_format = "0.0000"
        if r["Basis"] not in EQUITY_BASES:
            for c in row:
                c.fill = FALLBACK
    ws.auto_filter.ref = ws.dimensions

    # Monthly ----------------------------------------------------------------
    ws = wb.create_sheet("Monthly")
    _header(ws, ["Year", "Month", "Reporting days", "Layout", "Sum of daily Net (Rs Cr)",
                 "Sum of daily Net (US$ mn)", "NSDL month total Net (Rs Cr)",
                 "NSDL month total Net (US$ mn)", "Diff Rs Cr", "Diff US$ mn", "OK"],
            {4: 9, 5: 16, 6: 16, 7: 18, 8: 18})
    for r in reports:
        di, du, ok = check(r)
        st = r.stated or (None,) * 4
        ws.append([r.year, MONTHS[r.month - 1][:3], len(r.days), r.layout,
                   r.computed_net_inr, r.computed_net_usd, st[2], st[3], di, du,
                   "yes" if ok else ("fallback" if r.fallback_days else "NO")])
        row = ws[ws.max_row]
        for c in row[4:10]:
            c.number_format = DEC
        if not ok:
            for c in row:
                c.fill = FALLBACK if r.fallback_days else BAD
    ws.auto_filter.ref = ws.dimensions

    # Matrix -----------------------------------------------------------------
    ws = wb.create_sheet("INR_by_Year")
    _header(ws, ["Year"] + [m[:3] for m in MONTHS] + ["Year total"], {1: 8})
    years = sorted({r.year for r in reports})
    by = {(r.year, r.month): r for r in reports}
    for y in years:
        vals = [round(sum(d.net_inr or 0 for d in by[y, m].days), 2) if (y, m) in by else None
                for m in range(1, 13)]
        ws.append([y] + vals + [round(sum(v for v in vals if v is not None), 2)])
        row = ws[ws.max_row]
        for c in row[1:]:
            c.number_format = NUM
        row[13].font = Font(bold=True)
    ws.cell(row=len(years) + 3, column=1,
            value="Rs Crore, sum of the daily equity rows (fallback days count their Total row).").font = Font(italic=True)
    chart = BarChart()
    chart.title = "FPI net equity flow per year (Rs Crore)"
    chart.legend = None
    chart.height, chart.width = 9, 22
    chart.add_data(Reference(ws, min_col=14, min_row=1, max_row=len(years) + 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(years) + 1))
    ws.add_chart(chart, "P2")

    # Routes (daily, wide, two-row header) -------------------------------------
    ws = wb.create_sheet("Equity_Routes")
    short = [m.replace("Net Investment", "Net").replace("Crore", "Cr") for m in MEASURES]
    ws.append(["Date", "Year", "Month"] + [n for n in ROUTE_NAMES for _ in MEASURES] + ["USD/INR Rate"])
    ws.append(["", "", ""] + short * len(ROUTE_NAMES) + [""])
    for col in (1, 2, 3, 4 + len(ROUTE_NAMES) * len(MEASURES)):
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    for k in range(len(ROUTE_NAMES)):
        c0 = 4 + k * len(MEASURES)
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + len(MEASURES) - 1)
    for row in ws.iter_rows(min_row=1, max_row=2):
        for c in row:
            c.font, c.fill = HEAD, FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30
    for i in range(1, len(ROUTE_COLS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12 if i <= 3 else 14
    ws.freeze_panes = "D3"
    for r in rroutes:
        ws.append([r[c] for c in ROUTE_COLS])
        row = ws[ws.max_row]
        row[0].number_format = "dd-mmm-yyyy"
        for c in row[3:-1]:
            c.number_format = DEC
        row[-1].number_format = "0.0000"
    if rroutes:
        ws.auto_filter.ref = f"A2:{get_column_letter(len(ROUTE_COLS))}{ws.max_row}"

    # Routes (monthly reconciliation) ----------------------------------------
    ws = wb.create_sheet("Routes_Monthly")
    cols = ["Year", "Month"]
    for n in ROUTE_NAMES:
        cols += [f"{n}: sum of daily Net (Rs Cr)", f"{n}: NSDL month Net (Rs Cr)", f"{n}: diff"]
    _header(ws, cols + ["Max |diff| (all 4 columns)", "OK"], {i: 15 for i in range(3, len(cols) + 3)})
    ws.row_dimensions[1].height = 45
    for r in reports:
        worst, ok = check_routes(r)
        if worst is None and ok:
            continue                     # flat layout - no routes published
        vals = []
        for n in ROUTE_NAMES:
            st = r.stated_routes.get(n, (None,) * 4)[2]
            calc = r.route_sum(n)
            vals += [calc, st, None if st is None else round(calc - st, 2)]
        ws.append([r.year, MONTHS[r.month - 1][:3]] + vals + [worst, "yes" if ok else "NO"])
        row = ws[ws.max_row]
        for c in row[2:-1]:
            c.number_format = DEC
        if not ok:
            for c in row:
                c.fill = BAD
    ws.auto_filter.ref = ws.dimensions

    # Fallbacks --------------------------------------------------------------
    ws = wb.create_sheet("Fallbacks")
    _header(ws, DAILY_COLS, {1: 12, 4: 24})
    for r in fallbacks:
        ws.append([r[c] for c in DAILY_COLS])
        ws.cell(row=ws.max_row, column=1).number_format = "dd-mmm-yyyy"
    if not fallbacks:
        ws.append(["None - every reporting day in range has an Equity row."])

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
