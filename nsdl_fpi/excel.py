"""Workbook (and CSV) output for the FPI equity series."""
from __future__ import annotations

import csv
import os
from datetime import date, datetime, timezone

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .fetch import URL
from .parse import MONTHS, YearReport

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")
BAD = PatternFill("solid", fgColor="F8CBAD")
NUM = "#,##0;[Red]-#,##0"
COL = {"INR": "Equity (INR Crores)", "USD": "Equity (USD Million)"}


def _header(ws, cols, widths=None):
    ws.append(cols)
    for i, _ in enumerate(cols, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = HEAD, FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 14)
    ws.freeze_panes = "A2"


def long_rows(reports: list[YearReport], currencies: list[str]) -> list[dict]:
    by = {(r.year, r.currency): r for r in reports}
    years = sorted({r.year for r in reports})
    out = []
    for y in years:
        months = sorted({m for c in currencies if (y, c) in by for m in by[y, c].equity})
        for m in months:
            row = {"Period": date(y, m, 1), "Year": y, "Month": m, "Month Name": MONTHS[m - 1]}
            for c in currencies:
                row[COL[c]] = by[y, c].equity.get(m) if (y, c) in by else None
            out.append(row)
    return out


def write_csv(path: str, rows: list[dict]) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        for r in rows:
            w.writerow({**r, "Period": r["Period"].isoformat()})


def write_workbook(path: str, reports: list[YearReport], currencies: list[str]) -> None:
    rows = long_rows(reports, currencies)
    wb = Workbook()

    # README -----------------------------------------------------------------
    ws = wb.active
    ws.title = "README"
    first, last = rows[0]["Period"], rows[-1]["Period"]
    lines = [
        ("FPI Net Investment in Equity - monthly", True),
        ("", False),
        (f"Source: {URL}", False),
        ("Report: NSDL 'Monthly FPI Net Investments' (year-wise), Equity column only.", False),
        (f"Coverage: {first:%B %Y} to {last:%B %Y} ({len(rows)} months).", False),
        (f"Scraped: {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC", False),
        ("", False),
        ("Units: INR Crores and USD Million, exactly as NSDL publishes them (net = purchases - sales).", False),
        ("The FPI 'Equity' column is the first column of the report - NOT the 'Equity' sub-column", False),
        ("under Mutual Funds, which recent years also carry.", False),
        ("The current year is partial: only months NSDL has published are included.", False),
        ("NSDL compiles these from custodian reports; figures for recent months can be revised.", False),
        ("", False),
        ("Sheets", True),
        ("Equity_Monthly   one row per month, both currencies side by side (pivot-ready)", False),
        ("INR_by_Year      years down, months across, with NSDL's stated annual total", False),
        ("USD_by_Year      the same in USD Million", False),
        ("Check            sum of months vs NSDL's own 'Total - YYYY' row, per year and currency", False),
    ]
    for text, bold in lines:
        ws.append([text])
        if bold:
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=13 if ws.max_row == 1 else 11)
    ws.column_dimensions["A"].width = 100

    # Long sheet -------------------------------------------------------------
    ws = wb.create_sheet("Equity_Monthly")
    cols = list(rows[0])
    _header(ws, cols, {1: 12, 4: 12, 5: 20, 6: 20})
    for r in rows:
        ws.append([r[c] for c in cols])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "mmm-yyyy"
        for c in row[4:]:
            c.number_format = NUM
    ws.auto_filter.ref = ws.dimensions

    # Matrix sheets ----------------------------------------------------------
    by = {(r.year, r.currency): r for r in reports}
    for cur in currencies:
        ws = wb.create_sheet(f"{cur}_by_Year")
        _header(ws, ["Year"] + [m[:3] for m in MONTHS] + ["Total (NSDL)"],
                {1: 8, 14: 14})
        years = sorted(y for (y, c) in by if c == cur)
        for y in years:
            rep = by[y, cur]
            ws.append([y] + [rep.equity.get(m) for m in range(1, 13)] + [rep.stated_total])
        for row in ws.iter_rows(min_row=2):
            for c in row[1:]:
                c.number_format = NUM
            row[13].font = Font(bold=True)
        ws.cell(row=len(years) + 3, column=1, value=f"Units: {COL[cur]}").font = Font(italic=True)

        chart = BarChart()
        chart.title = f"Annual FPI net equity flow ({COL[cur].split('(')[1].rstrip(')')})"
        chart.y_axis.title = COL[cur].split("(")[1].rstrip(")")
        chart.legend = None
        chart.height, chart.width = 9, 22
        chart.add_data(Reference(ws, min_col=14, min_row=1, max_row=len(years) + 1), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(years) + 1))
        ws.add_chart(chart, "P2")

    # Check sheet ------------------------------------------------------------
    ws = wb.create_sheet("Check")
    _header(ws, ["Year", "Currency", "Months", "Sum of months", "NSDL stated total", "Difference", "OK"],
            {4: 16, 5: 18})
    for rep in sorted(reports, key=lambda r: (r.currency != "INR", r.year)):
        comp, st = rep.computed_total, rep.stated_total
        diff = None if comp is None or st is None else round(comp - st, 2)
        ok = diff is not None and abs(diff) <= 1      # NSDL rounds each month
        ws.append([rep.year, rep.currency, len(rep.equity), comp, st, diff, "yes" if ok else "NO"])
        if not ok:
            for c in ws[ws.max_row]:
                c.fill = BAD
    for row in ws.iter_rows(min_row=2):
        for c in row[3:6]:
            c.number_format = NUM

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


def mismatches(reports: list[YearReport]) -> list[YearReport]:
    out = []
    for r in reports:
        c, s = r.computed_total, r.stated_total
        if c is None or s is None or abs(c - s) > 1:
            out.append(r)
    return out
