"""APMI IA Insights: strategy, inception date and service type per approach.

Source: <https://insights.apmiindia.org> (APMI's "IA Insight Report" site).

How the data is fetched
-----------------------
The report page (``/ia-insight-report/<id>``) is rendered in the browser from
Next.js *server actions*, so its HTML holds no data. Three actions are used:

``searchAPMI("%")``     the site's search box; ``%`` is a SQL wildcard, so this
                        returns every approach APMI holds (id, name, manager) -
                        including dormant ones the ranking pages leave out
``fetchByServiceType``  the "Top (Non-)Discretionary Approaches" ranking pages:
                        strategy, service type and inception date for every
                        listed approach, 500 per call
``fetchIaDetails``      the report page's own header (the same three fields),
                        fetched only for matched approaches the rankings omit

Action ids change when APMI redeploys; :meth:`ApmiClient.call` re-reads them
from the page's JavaScript when a known id stops working.

How approaches are matched to the workbook
------------------------------------------
The workbook's ``Unique`` sheet lists approach names only. Tables B and C give
each discretionary name its manager's SEBI registration number(s); APMI gives a
manager *name*, which is mapped onto registration numbers through the
``Managers`` sheet. A match needs the name to agree **and** the manager to
agree, tried in three passes:

``Exact``       names equal ignoring case, spaces and punctuation
``Normalised``  also ignoring generic words (Approach, Strategy, PMS, Fund ...)
``Fuzzy``       same manager, names differ only by typos: each word the same
                or >= 80 % alike, extra words only initials or the manager's
                name, identical numbers / roman numerals and "non"/"ND"
                markers; or identical once the manager's brand is dropped
                ("Scient Smart Beta PMS") - always flagged for review

Non-discretionary names are not in tables B/C (table G has no approaches), so
they carry no registration number and match on name alone ("name only").
"""
from __future__ import annotations

import difflib
import json
import logging
import os
import re
import time
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime

import requests

log = logging.getLogger(__name__)

BASE = "https://insights.apmiindia.org"
LIST_PAGE = BASE + "/ia-by-service-type/1"
REPORT_URL = BASE + "/ia-insight-report/{}"
#: Any report page serves the report-page actions; this one is APMI's own example.
REPORT_PAGE = REPORT_URL.format("019936d5-e4e8-72e6-b1a9-5651e7b251a7")

#: action name -> (page that hosts it, id as deployed in Sep-2026)
ACTIONS = {
    "fetchByServiceType": (LIST_PAGE, "7fbcbf717c7c6cc58586e4f569087f857eba792dbe"),
    "searchAPMI": (REPORT_PAGE, "7f858098083f20a3b2aff06de708558059f200fcc1"),
    "fetchIaDetails": (REPORT_PAGE, "7f47cca4a0daab30bb6d23780660989ac78b80f04f"),
}
#: APMI's service-type ids for the ranking call.
SERVICE_TYPES = (1, 2)
SERVICE_LABEL = {"D": "Discretionary", "N": "Non-Discretionary"}

UA = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) "
      "Chrome/124.0 Safari/537.36 sebi-pmr-scraper/1.0")

_REF = re.compile(r'createServerReference\)\("([0-9a-f]{40,})"[^)]*?"([A-Za-z_]\w*)"\)')
_CHUNK = re.compile(r'(/_next/static/chunks/[^"\\\s]+\.js)')


class NoData(Exception):
    """APMI answered, but with an error instead of data."""


def parse_rsc(text: str):
    """The value a server action returned (row ``1:`` of the RSC stream)."""
    for line in text.splitlines():
        if line.startswith("1:E"):
            raise NoData(line[3:])
        if line.startswith("1:"):
            return json.loads(line[2:])
    raise NoData(text[:200])


class ApmiClient:
    def __init__(self, timeout: float = 60, delay: float = 0.5,
                 session: requests.Session | None = None):
        self.s = session or requests.Session()
        self.s.headers["User-Agent"] = UA
        self.timeout, self.delay = timeout, delay
        self.ids = {name: known for name, (_, known) in ACTIONS.items()}
        self._rediscovered: set[str] = set()

    def discover(self, name: str) -> str:
        """Read an action's id out of its page's scripts."""
        page = ACTIONS[name][0]
        html = self.s.get(page, timeout=self.timeout).text
        for path in dict.fromkeys(_CHUNK.findall(html)):
            js = self.s.get(BASE + path, timeout=self.timeout).text
            for action, found in _REF.findall(js):
                if found == name:
                    return action
        raise RuntimeError(f"could not find the {name} action on {page}")

    def call(self, name: str, args: list, retries: int = 3):
        page = ACTIONS[name][0]
        for attempt in range(retries):
            try:
                r = self.s.post(page, data=json.dumps(args), timeout=self.timeout,
                                headers={"Next-Action": self.ids[name],
                                         "Accept": "text/x-component",
                                         "Content-Type": "text/plain;charset=UTF-8"})
            except requests.RequestException as e:
                log.warning("%s %s: %s", name, args, e)
                time.sleep(2 ** attempt)
                continue
            finally:
                time.sleep(self.delay)
            stale = r.status_code == 404 or (
                r.ok and "text/x-component" not in r.headers.get("content-type", ""))
            if stale and name not in self._rediscovered:
                # an unknown action id gets a 404 or an HTML page, not an RSC stream
                self._rediscovered.add(name)
                self.ids[name] = self.discover(name)
                log.info("%s action id changed; now %s", name, self.ids[name])
                continue
            if "\n1:E" in r.text:
                # the action ran and failed on APMI's side for this record;
                # it fails the same way every time, so do not retry
                return parse_rsc(r.text)
            if r.status_code >= 500 or r.status_code == 429:
                time.sleep(2 ** attempt)
                continue
            r.raise_for_status()
            return parse_rsc(r.text)
        raise NoData(f"{name} {args}: no answer after {retries} tries")

    def search_all(self) -> list[dict]:
        """Every approach APMI holds: ``[{id, iaName, pmsName}]``."""
        rows = self.call("searchAPMI", ["%"])
        return [{"id": r["iaId"], "iaName": r["iaName"], "pmsName": r["pmsName"]} for r in rows]

    def rankings(self, page_size: int = 500) -> dict[str, dict]:
        """Details of every approach on the ranking pages, by id."""
        out = {}
        for st in SERVICE_TYPES:
            page = 1
            while True:
                # service type, sort (3 = site default), age filter, AUM filter, size, page
                batch = self.call("fetchByServiceType",
                                  [st, 3, "$undefined", "$undefined", page_size, page])
                for a in batch:
                    out[a["pmsIaId"]] = {
                        "id": a["pmsIaId"], "iaName": a["iaName"], "pmsName": a["pmsProviderName"],
                        "strategy": a["iaStrategy"], "serviceType": a["iaServiceType"],
                        "dateOfInception": a["iaDateOfInception"], "source": "ranking",
                    }
                if len(batch) < page_size:
                    break
                page += 1
        return out

    def details(self, ia_id: str) -> dict:
        a = self.call("fetchIaDetails", [ia_id])
        return {"id": a["id"], "iaName": a["iaName"], "pmsName": a["pmsName"],
                "strategy": a["strategy"], "serviceType": a["serviceType"],
                "dateOfInception": a["dateOfInception"], "source": "report"}


# --------------------------------------------------------------------------
# name normalisation

_LEGAL = re.compile(r"\b(private|pvt|limited|ltd|lmited|lted|llp|the|co|company|"
                    r"corporation|corp|india|p|ms)\b")
_GENERIC = re.compile(r"\b(approach|strategy|portfolio|pms|dpms|fund|plan|scheme|"
                      r"ia|investment|the|of|and|option|series)\b")
_ROMAN = re.compile(r"^(?=[ivxl]+$)(xc|xl|l?x{0,3})(ix|iv|v?i{0,3})$")
_NEGATION = {"non", "nd", "ndpms"}


def manager_key(name: str) -> str:
    s = str(name).lower().replace("&", " and ")
    s = s.replace("limited liability partnership", "llp")
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    return re.sub(r"\s+", "", _LEGAL.sub(" ", s))


def manager_variants(name: str) -> set[str]:
    """Keys for a manager name, with and without "(Formerly known as ...)"."""
    bare = re.sub(r"\([^)]*\)", "", name)
    out = {name, bare, re.sub(r"\s*-\s*portfolio managers?\s*$", "", bare, flags=re.I)}
    for inner in re.findall(r"\(([^)]*)\)", name):
        out.add(re.sub(r"^(formerly known as|formerly|earlier known as)\s*", "",
                       inner.strip(), flags=re.I))
    return {k for k in map(manager_key, out) if k}


def exact_key(name) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


def loose_key(name) -> str:
    s = str(name).lower().replace("&", " and ")
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    return re.sub(r"\s+", "", _GENERIC.sub(" ", s))


def _markers(name) -> tuple:
    toks = re.findall(r"[a-z]+|\d+", str(name).lower())
    nums = sorted(t for t in toks if t.isdigit() or _ROMAN.match(t))
    return nums, sorted(t for t in toks if t in _NEGATION)


def _words(name) -> list[str]:
    return [t for t in re.findall(r"[a-z0-9]+", str(name).lower().replace("&", " and "))
            if not _GENERIC.fullmatch(t)]


def _initials(words: list[str]) -> str:
    return "".join(w[0] for w in words if w)


def _is_subsequence(short: str, long: str) -> bool:
    it = iter(long)
    return all(ch in it for ch in short)


def _ignorable(word: str, name: str, manager: str | None) -> bool:
    """Words one name may carry and the other not: initials, acronyms, the
    manager's own name ("RH", "RHPMPL", "BB", "MOWPL", "MMFS", "Val-Q", "io")."""
    if len(word) <= 2 or word == "io":
        return True
    raw = re.findall(r"[a-z0-9]+", str(name).lower())
    if _initials([w for w in raw if w != word]) == word:
        return True
    if manager:
        mwords = re.findall(r"[a-z0-9]+", manager.lower().replace("&", " and "))
        inits = _initials(mwords)
        if word in mwords or (len(word) >= 2 and (_is_subsequence(word, inits)
                                                  or _is_subsequence(inits, word))):
            return True
    return False


def _without_manager(name, manager: str) -> tuple:
    """Loose key and markers of ``name`` once words of the manager's name are gone."""
    mwords = set(re.findall(r"[a-z0-9]+", manager.lower()))
    kept = " ".join(w for w in re.findall(r"[a-z0-9]+", str(name).lower().replace("&", " and "))
                    if w not in mwords)
    key = loose_key(kept)
    return (key, str(_markers(kept))) if len(key) >= 4 else ("", "")


def fuzzy_score(a, b, manager: str | None = None, cutoff: float = 0.85) -> float:
    """Similarity if names ``a`` and ``b`` differ only by typos, else 0.

    Needs identical numbers / roman numerals and "non"/"ND" markers, every word
    of one name to be the same as or a typo of a word of the other (>= 80 %
    alike), and any word left over to be an initialism or the manager's name.
    """
    if manager and _without_manager(a, manager) == _without_manager(b, manager) != ("", ""):
        return 0.9       # same once the manager's brand is dropped ("Scient Smart Beta PMS")
    if _markers(a) != _markers(b):
        return 0.0
    ratio = difflib.SequenceMatcher(None, exact_key(a), exact_key(b)).ratio()
    if ratio < cutoff:
        return 0.0
    wa, wb = _words(a), _words(b)
    left_b = list(wb)
    left_a = []
    for w in wa:
        best = max(left_b, key=lambda x: difflib.SequenceMatcher(None, w, x).ratio(), default=None)
        if best is not None and difflib.SequenceMatcher(None, w, best).ratio() >= 0.8:
            left_b.remove(best)
        else:
            left_a.append(w)
    left_a = [w for w in left_a if not _ignorable(w, a, manager)]
    left_b = [w for w in left_b if not _ignorable(w, b, manager)]
    if left_a or left_b:
        # words run together or split apart ("EMKAYS GOLDEN" / "EMKAYGOLDEN")
        joined = difflib.SequenceMatcher(None, "".join(left_a), "".join(left_b)).ratio()
        if joined < 0.8:
            return 0.0
    return ratio


# --------------------------------------------------------------------------
# matching

@dataclass
class Match:
    approach: str
    reg_no: str | None
    manager: str | None
    how: str
    apmi: dict | None = None
    note: str = ""


def provider_regs(names: set[str], managers: dict[str, list[str]]) -> dict[str, set[str]]:
    """APMI manager name -> SEBI registration number(s).

    ``managers`` maps registration number -> every name SEBI printed for it.
    One name can hold two registrations (managers that re-registered), so the
    result is a set; it is empty when nothing on SEBI's side resembles it.
    """
    by_key: dict[str, set[str]] = defaultdict(set)
    for reg, ns in managers.items():
        for n in ns:
            by_key[manager_key(n)].add(reg)
    keys = list(by_key)
    out = {}
    for prov in names:
        regs: set[str] = set()
        for v in manager_variants(prov):
            regs |= by_key.get(v, set())
        if not regs:
            for v in manager_variants(prov):
                for k in difflib.get_close_matches(v, keys, n=1, cutoff=0.88):
                    regs |= by_key[k]
        out[prov] = regs
    return out


PLACEHOLDERS = {"0", "NA"}
#: APMI's own test entries, not real managers.
TEST_MANAGERS = {"apmitest"}


def match_all(unique: list[str], regs_by_name: dict[str, set[str]],
              reg_names: dict[str, str], apmi: list[dict],
              prov_regs: dict[str, set[str]], ranked: set[str] = frozenset()) -> list[Match]:
    """One :class:`Match` per (approach name, manager) - more when APMI has
    several equally good candidates, all of which are listed and flagged.

    Exact and normalised matches are made for every name first, so a fuzzy
    match can say which other name on the Unique sheet already matches its
    APMI approach exactly (usually an older or corrected spelling of it).
    ``ranked`` (ids on APMI's ranking pages) breaks ties between APMI's own
    duplicate entries for one manager.
    """
    apmi = [a for a in apmi if manager_key(a["pmsName"]) not in TEST_MANAGERS]
    by_exact, by_loose, by_reg = defaultdict(list), defaultdict(list), defaultdict(list)
    for a in apmi:
        by_exact[exact_key(a["iaName"])].append(a)
        if len(loose_key(a["iaName"])) >= 3:
            by_loose[loose_key(a["iaName"])].append(a)
        for r in prov_regs.get(a["pmsName"], ()):
            by_reg[r].append(a)

    def same_mgr(a, reg):
        return reg in prov_regs.get(a["pmsName"], ())

    pending: list[tuple] = []           # (name, reg) awaiting the fuzzy pass
    found: dict[tuple, tuple] = {}
    for name in unique:
        if str(name).strip().upper() in PLACEHOLDERS:
            continue
        lk = loose_key(name)
        for reg in sorted(regs_by_name.get(name, ())) or [None]:
            hit = None
            for how, cands in (("Exact", by_exact.get(exact_key(name), [])),
                               ("Normalised", by_loose.get(lk, []) if len(lk) >= 3 else [])):
                if how == "Normalised":
                    cands = [a for a in cands if _markers(a["iaName"]) == _markers(name)]
                if reg is not None:
                    cands = [a for a in cands if same_mgr(a, reg)]
                if cands:
                    hit = (how + (" (name + manager)" if reg else " (name only)"),
                           _best(name, cands, ranked), "")
                    break
            if hit:
                found[(name, reg)] = hit
            else:
                pending.append((name, reg))

    claimed: dict[str, list[str]] = defaultdict(list)
    for (name, _), (_, cands, _) in found.items():
        for a in cands:
            claimed[a["id"]].append(name)
    for name, reg in pending:
        pool = by_reg.get(reg, ()) if reg else ()
        scored = sorted(((fuzzy_score(name, a["iaName"], a["pmsName"]), a) for a in pool),
                        key=lambda x: -x[0])
        if scored and scored[0][0] and (len(scored) == 1 or scored[1][0] < scored[0][0]):
            best, a = scored[0]
            note = f"Names {best:.0%} alike - confirm it is the same approach"
            others = sorted(set(claimed.get(a["id"], ())))
            if others:
                note += "; APMI's approach also matches " + ", ".join(f'"{o}"' for o in others[:3])
            found[(name, reg)] = ("Fuzzy (same manager) - check", [a], note)

    out: list[Match] = []
    for name in unique:
        if str(name).strip().upper() in PLACEHOLDERS:
            out.append(Match(name, None, None, "SEBI placeholder",
                             note="Not an approach: SEBI's placeholder for a manager "
                                  "with no discretionary approaches"))
            continue
        for reg in sorted(regs_by_name.get(name, ())) or [None]:
            mgr = reg_names.get(reg) if reg else None
            if (name, reg) not in found:
                out.append(Match(name, reg, mgr, "Not on APMI"))
                continue
            how, cands, note = found[(name, reg)]
            if len(cands) > 1:
                note = f"{len(cands)} APMI approaches fit - all listed, check which applies"
            for a in cands:
                out.append(Match(name, reg, mgr, how, a, note))
    return out


def _best(name, cands: list[dict], ranked: set[str] = frozenset()) -> list[dict]:
    """Narrow several candidates: same ND/non markers, then identical spelling,
    then - only among one manager's duplicates - the entry APMI ranks."""
    one_manager = len({manager_key(a["pmsName"]) for a in cands}) == 1
    for keep in (lambda a: _markers(a["iaName"]) == _markers(name),
                 lambda a: re.sub(r"\s+", "", a["iaName"].lower())
                 == re.sub(r"\s+", "", str(name).lower()),
                 lambda a: one_manager and a["id"] in ranked):
        if len(cands) > 1:
            narrowed = [a for a in cands if keep(a)]
            cands = narrowed or cands
    return cands


# --------------------------------------------------------------------------
# fetch + cache

def load_apmi(client: ApmiClient, cache_path: str | None = None, refresh: bool = False) -> dict:
    """Every APMI approach (``universe``) and the ranking pages' details.

    Cached in ``cache_path`` with the report-page details :func:`fill_details`
    adds, so a re-run asks APMI only for what it lacks; ``refresh`` starts over.
    """
    cache = {}
    if cache_path and os.path.exists(cache_path) and not refresh:
        with open(cache_path) as f:
            cache = json.load(f)
    if "universe" not in cache:
        cache["universe"] = client.search_all()
        cache["rankings"] = client.rankings()
        cache["reports"] = {}
        cache["fetched_at"] = datetime.now().strftime("%d-%b-%Y")
    return cache


def fill_details(client: ApmiClient, cache: dict, wanted: set[str],
                 cache_path: str | None = None) -> dict[str, dict | None]:
    """Details for each id in ``wanted``: ranking pages first, else its report page."""
    todo = [i for i in sorted(wanted)
            if i not in cache["rankings"] and i not in cache["reports"]]
    log.info("%d approaches need their report page", len(todo))
    for n, ia_id in enumerate(todo, 1):
        try:
            cache["reports"][ia_id] = client.details(ia_id)
        except NoData as e:
            log.info("no report data for %s: %s", ia_id, e)
            cache["reports"][ia_id] = None
        if cache_path and (n % 50 == 0 or n == len(todo)):
            _save(cache, cache_path)
            log.info("report pages: %d / %d", n, len(todo))
    if cache_path:
        _save(cache, cache_path)
    return {i: cache["rankings"].get(i) or cache["reports"].get(i) for i in wanted}


def _save(cache, path):
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w") as f:
        json.dump(cache, f)
    os.replace(tmp, path)


# --------------------------------------------------------------------------
# workbook

def read_workbook_inputs(path: str):
    """Approach names, their managers and every manager name, from the workbook."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    unique = [r[1] for r in list(wb["Unique"].iter_rows(values_only=True))[2:]
              if r[1] is not None]
    regs_by_name: dict[str, set[str]] = defaultdict(set)
    for sheet in ("B. AUM (Discretionary)", "C. Flows (Discretionary)"):
        for r in wb[sheet].iter_rows(min_row=5, values_only=True):
            if r[0] and r[2] and r[3] is not None:
                regs_by_name[r[3]].add(r[2])
    managers, reg_names = {}, {}
    for r in list(wb["Managers"].iter_rows(values_only=True))[4:]:
        if not r[0]:
            continue
        reg_names[r[0]] = r[1]
        others = [re.sub(r"\s*\([A-Z][a-z]{2}-\d{4} to [A-Z][a-z]{2}-\d{4}\)\s*$", "", p).strip()
                  for p in re.split(r";\s*", str(r[5] or ""))]
        managers[r[0]] = [r[1]] + [o for o in others if o]
    wb.close()
    return unique, regs_by_name, managers, reg_names


def inception(value: str | None):
    try:
        return datetime.strptime(value, "%d-%m-%Y")
    except (TypeError, ValueError):
        return None


COLUMNS = [  # header, width
    ("Investment Approach", 48), ("Portfolio Manager", 40), ("Registration No.", 15),
    ("Strategy", 12), ("Discretionary / Non-Discretionary", 20), ("Date of Inception", 13),
    ("Name on APMI", 48), ("Portfolio Manager on APMI", 40), ("Match", 30),
    ("IA Insight Report", 14), ("Note", 60),
]


def write_sheet(wb, rows: list[Match], details: dict[str, dict | None],
                fetched_at: str, title: str = "IA Insights"):
    """Add (or replace) the IA Insights sheet, styled like the workbook's own."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    navy, body = "FF154063", "FF333333"
    edge = Side(style="thin", color="FFBBBBBB")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    stripes = [PatternFill("solid", fgColor="FFECECEC"), PatternFill("solid", fgColor="FFFFFFFF")]

    ws["A1"] = "IA Insights"
    ws["A1"].font = Font(name="Roboto", size=10.5, bold=True, color=navy)
    ws["A2"] = ("Strategy, date of inception and discretionary / non-discretionary tag for "
                "every approach on the Unique sheet, from APMI's IA Insight Report "
                f"(insights.apmiindia.org), fetched {fetched_at}. Blank = not found on APMI; "
                "the Match column says how each row was matched.")
    ws["A2"].font = Font(name="Roboto", size=9.5, color="FF666666")

    for c, (head, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(4, c, head)
        cell.font = Font(name="Roboto", size=10, bold=True, color=navy)
        cell.fill = stripes[1]
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[4].height = 28

    font = Font(name="Roboto", size=9, color=body)
    link_font = Font(name="Roboto", size=9, color="FF0563C1", underline="single")
    centre = Alignment(horizontal="center")
    for i, m in enumerate(rows):
        r = 5 + i
        a = m.apmi or {}
        d = details.get(a["id"]) if a else None
        note = m.note
        if a and d is None:
            note = "; ".join(filter(None, [note, "APMI's report page for this approach "
                                                 "shows \"Failed to load data\""]))
        values = [m.approach, m.manager or a.get("pmsName"), m.reg_no,
                  d and d["strategy"], d and SERVICE_LABEL.get(d["serviceType"], d["serviceType"]),
                  d and inception(d["dateOfInception"]),
                  a.get("iaName"), a.get("pmsName"), m.how,
                  "View report" if a else None, note or None]
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            cell.font = font
            cell.fill = stripes[i % 2]
            cell.border = border
        for c in (3, 4, 5, 6):
            ws.cell(r, c).alignment = centre
        ws.cell(r, 6).number_format = "DD-MM-YYYY"
        if a:
            link = ws.cell(r, 10)
            link.hyperlink = REPORT_URL.format(a["id"])
            link.font = link_font
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(COLUMNS))}{4 + len(rows)}"
    return ws


def append_rows(ws, rows: list[tuple]):
    """Append rows to a Notes/Verification-style sheet, continuing its striping."""
    from copy import copy
    last = ws.max_row
    while last > 4 and all(ws.cell(last, c).value is None for c in range(1, ws.max_column + 1)):
        last -= 1
    for vals in rows:
        last += 1
        like = ws.cell(last - 2, 1).row if last - 2 >= 5 else last - 1  # same stripe
        for c, v in enumerate(vals, 1):
            cell, src = ws.cell(last, c, v), ws.cell(like, c)
            cell.font, cell.fill = copy(src.font), copy(src.fill)
            cell.border, cell.alignment = copy(src.border), copy(src.alignment)
    if ws.auto_filter.ref:
        start, end = ws.auto_filter.ref.split(":")
        end = re.sub(r"\d+$", str(last), end)
        ws.auto_filter.ref = f"{start}:{end}"


def spot_check(client: ApmiClient, details: dict[str, dict | None], n: int = 25,
               seed: int = 2026) -> tuple[int, list[str]]:
    """Re-read ``n`` ranking-page approaches from their report pages and compare."""
    import random
    ranked = sorted(i for i, d in details.items() if d and d["source"] == "ranking")
    sample = random.Random(seed).sample(ranked, min(n, len(ranked)))
    bad = []
    for ia_id in sample:
        want = details[ia_id]
        try:
            got = client.details(ia_id)
        except NoData as e:
            bad.append(f"{want['iaName']}: report page gave no data ({e})")
            continue
        diff = [k for k in ("strategy", "serviceType", "dateOfInception") if got[k] != want[k]]
        if diff:
            bad.append(f"{want['iaName']}: " + ", ".join(
                f"{k} {want[k]!r} vs {got[k]!r}" for k in diff))
    return len(sample), bad


@dataclass
class Summary:
    names: int
    rows: int
    by_how: dict
    names_found: int
    with_details: int
    no_details: int
    spot: tuple[int, list[str]] | None = None

    def text(self) -> str:
        parts = ", ".join(f"{k}: {v:,}" for k, v in sorted(self.by_how.items(), key=lambda x: -x[1]))
        return (f"{self.names:,} approach names -> {self.rows:,} rows; {self.names_found:,} "
                f"names found on APMI ({parts}); {self.no_details:,} matched rows without "
                f"APMI data")

    def verification_rows(self) -> list[tuple]:
        rows = [
            ("IA Insights: every approach on the Unique sheet has a row",
             "PASS" if self.names else "FAIL",
             f"{self.names:,} names -> {self.rows:,} rows (one per manager using the name, "
             f"more where APMI has several candidates)"),
            ("IA Insights: matches", "INFO",
             "; ".join(f"{k} {v:,}" for k, v in sorted(self.by_how.items(), key=lambda x: -x[1]))),
            ("IA Insights: every matched row has strategy, service type and inception date",
             "PASS" if not self.no_details else "INFO",
             f"{self.with_details:,} rows filled; {self.no_details:,} where APMI's own report "
             "page shows \"Failed to load data\" (noted on the row)"),
        ]
        if self.spot:
            n, bad = self.spot
            rows.append((
                "IA Insights: ranking-page figures equal the IA Insight Report page",
                "PASS" if not bad else "FAIL",
                f"{n} random approaches re-read from their report pages; "
                + ("all three fields identical" if not bad else "; ".join(bad))))
        return rows


def summarise(rows: list[Match], unique: list[str], details: dict[str, dict | None]) -> Summary:
    by_how: dict[str, int] = defaultdict(int)
    for m in rows:
        by_how[m.how] += 1
    matched = [m for m in rows if m.apmi]
    filled = sum(1 for m in matched if details.get(m.apmi["id"]))
    return Summary(
        names=len(unique), rows=len(rows), by_how=dict(by_how),
        names_found=len({m.approach for m in matched}),
        with_details=filled, no_details=len(matched) - filled,
    )


def notes_rows(fetched_at: str) -> list[tuple]:
    return [(
        "IA Insights",
        "Strategy (Equity / Debt / Hybrid / Multi Asset), Discretionary / Non-Discretionary and "
        "Date of Inception for every name on the Unique sheet, from APMI's IA Insight Report "
        f"(https://insights.apmiindia.org), fetched {fetched_at}. A name used by several managers "
        "gets one row per manager (manager and registration number from tables B and C). Match: "
        "Exact = same name and same manager; Normalised = same once generic words such as "
        "Approach / Strategy / PMS are ignored; Fuzzy = same manager, name differs only by typos "
        "(check these); name only = non-discretionary names, which SEBI's tables do not tie to a "
        "manager; Not on APMI = APMI has no approach of that name for that manager (mostly closed "
        "or renamed approaches). View report opens the approach's APMI page.",
    )]
