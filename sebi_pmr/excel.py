"""Write the scraped tables into one organised workbook (plus CSV siblings).

Layout - one sheet per SEBI table, each row a (manager, month, approach) fact:

======================  ==========================================================
Sheet                   Contents
======================  ==========================================================
``README``              provenance, column glossary, the Feb-2021 format caveat
``B_Disc_AUM``          table B, every AUM column, discretionary
``C_Disc_NetFlow``      table C, net inflow/outflow during the month only
``G_NonDisc_AUM``       table G, every AUM column, non-discretionary
``H_NonDisc_Flow``      table H, the "during the month" inflow/outflow/net group
``Coverage``            per-month fetch outcomes, so gaps are visible not silent
======================  ==========================================================

Excel caps a sheet at 1 048 576 rows; a table larger than that is split into
``_pt2``, ``_pt3`` ... sheets rather than being silently truncated.
"""
from __future__ import annotations

import csv
import json
import logging
import os
from calendar import monthrange
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

EXCEL_MAX_ROWS = 1_048_576
SHEET_ROW_LIMIT = EXCEL_MAX_ROWS - 10

#: Preferred left-to-right order for the AUM break-up columns.
AUM_ORDER = [
    "Equity - Listed", "Equity - Unlisted",
    "Plain Debt - Listed", "Plain Debt - Unlisted",
    "Structured Debt - Listed", "Structured Debt - Unlisted",
    "Derivatives - Equity", "Derivatives - Commodity", "Derivatives - Others",
    "Mutual Funds", "Others", "Total",
]

MONTH_NAMES = ["", "January", "February", "March", "April", "May", "June",
               "July", "August", "September", "October", "November", "December"]

ID_COLS = ["Portfolio Manager", "Registration No", "Year", "Month", "Month Name", "Period"]

_HDR_FILL = PatternFill("solid", fgColor="1F3864")
_HDR_FONT = Font(color="FFFFFF", bold=True, size=10)


def month_end(year: int, month: int) -> date:
    return date(year, month, monthrange(year, month)[1])


def _aum_columns(rows: list[dict]) -> list[str]:
    """Canonical AUM columns first, then anything new SEBI has added."""
    seen = set()
    for r in rows:
        seen.update(r.keys())
    ordered = [c for c in AUM_ORDER if c in seen]
    return ordered + sorted(seen - set(ordered))


def _latest_names(conn) -> dict[str, tuple[str, str]]:
    """Most recent name/registration per manager.

    Managers rebrand (IIFL Asset Management -> 360 ONE Asset Management) while
    keeping one registration number, so the newest name is used throughout and
    the registration number stays the join key.
    """
    out = {}
    for r in conn.execute(
        "SELECT pmr_id, pm_name, reg_no FROM report_meta "
        "ORDER BY year DESC, month DESC"
    ):
        out.setdefault(r["pmr_id"], (r["pm_name"], r["reg_no"]))
    for r in conn.execute("SELECT pmr_id, name, reg_no FROM pm"):
        out.setdefault(r["pmr_id"], (r["name"], r["reg_no"]))
    return out


def collect(conn):
    """Read every table out of SQLite into plain row dicts."""
    names = _latest_names(conn)

    def ident(pmr_id, year, month):
        nm, reg = names.get(pmr_id, (pmr_id, ""))
        return [nm, reg, year, month, MONTH_NAMES[month], month_end(year, month)]

    b_raw = [(r["pmr_id"], r["year"], r["month"], r["investment_approach"],
              r["is_total"], json.loads(r["data"]))
             for r in conn.execute("SELECT * FROM b_aum ORDER BY year,month,pmr_id,investment_approach")]
    b_cols = _aum_columns([d for *_, d in b_raw])
    b_rows = [ident(p, y, m) + [a, bool(t)] + [d.get(c) for c in b_cols]
              for p, y, m, a, t, d in b_raw]

    c_rows = [ident(r["pmr_id"], r["year"], r["month"]) +
              [r["investment_approach"], bool(r["is_total"]), r["net_month"]]
              for r in conn.execute("SELECT * FROM c_flow ORDER BY year,month,pmr_id,investment_approach")]

    g_raw = [(r["pmr_id"], r["year"], r["month"], json.loads(r["data"]))
             for r in conn.execute("SELECT * FROM g_aum ORDER BY year,month,pmr_id")]
    g_cols = _aum_columns([d for *_, d in g_raw])
    g_rows = [ident(p, y, m) + [d.get(c) for c in g_cols] for p, y, m, d in g_raw]

    h_rows = [ident(r["pmr_id"], r["year"], r["month"]) +
              [r["inflow_month"], r["outflow_month"], r["net_month"]]
              for r in conn.execute("SELECT * FROM h_flow ORDER BY year,month,pmr_id")]

    quality = _quality(b_raw, g_raw, b_cols, g_cols, names)

    return {
        "B_Disc_AUM": (ID_COLS + ["Investment Approach", "Is Total Row"] + b_cols, b_rows),
        "C_Disc_NetFlow": (ID_COLS + ["Investment Approach", "Is Total Row",
                                      "Net Inflow(+)/Outflow(-) During Month (INR cr)"], c_rows),
        "G_NonDisc_AUM": (ID_COLS + g_cols, g_rows),
        "H_NonDisc_Flow": (ID_COLS + ["Inflow During Month (INR cr)",
                                      "Outflow During Month (INR cr)",
                                      "Net Inflow(+)/Outflow(-) During Month (INR cr)"], h_rows),
        "Data_Quality": quality,
    }


QUALITY_ROW_CAP = 50_000


def _quality(b_raw, g_raw, b_cols, g_cols, names, tolerance: float = 0.05):
    """Rows whose component columns do not add up to SEBI's stated Total.

    These are inconsistencies in the filings themselves, not extraction errors:
    the *column-wise* check (SEBI's Total row against the sum of its approach
    rows) reconciles exactly, while some individual approach rows as filed do
    not.  They are listed so nobody silently averages over a bad figure.
    """
    head = ["Sheet", "Portfolio Manager", "Registration No", "Period",
            "Investment Approach", "Sum of Components", "Stated Total", "Difference"]
    out = []
    for sheet, raw, cols, labelled in (("B_Disc_AUM", b_raw, b_cols, True),
                                       ("G_NonDisc_AUM", g_raw, g_cols, False)):
        comps = [c for c in cols if c != "Total"]
        for rec in raw:
            if labelled:
                pmr_id, year, month, approach, _is_total, data = rec
            else:
                pmr_id, year, month, data = rec
                approach = "(non-discretionary total)"
            if "Total" not in data:
                continue
            total = data.get("Total")
            if total is None:
                continue
            ssum = sum((data.get(c) or 0) for c in comps)
            if abs(ssum - total) > tolerance:
                nm, reg = names.get(pmr_id, (pmr_id, ""))
                out.append([sheet, nm, reg, month_end(year, month), approach,
                            round(ssum, 4), total, round(ssum - total, 4)])
    if len(out) > QUALITY_ROW_CAP:
        # A flood here means something systemic; the log is the signal, and the
        # sheet stays readable instead of dwarfing the data it annotates.
        log.warning("%d reconciliation mismatches found; listing the first %d",
                    len(out), QUALITY_ROW_CAP)
        out = out[:QUALITY_ROW_CAP]
    return head, out


def coverage_rows(conn):
    head = ["Year", "Month", "Month Name", "Managers Fetched",
            "With Data", "No Data Filed", "Legacy Format", "Errors"]
    out = []
    for r in conn.execute(
        "SELECT year, month, COUNT(*) n,"
        " SUM(status='ok') ok, SUM(status='nodata') nd,"
        " SUM(status='oldformat') old, SUM(status='error') err"
        " FROM page_log GROUP BY year, month ORDER BY year, month"
    ):
        out.append([r["year"], r["month"], MONTH_NAMES[r["month"]], r["n"],
                    r["ok"], r["nd"], r["old"], r["err"]])
    return head, out


def _header_cells(ws, head: list[str]):
    """Styled header row for a write-only sheet."""
    from openpyxl.cell import WriteOnlyCell
    out = []
    for name in head:
        cell = WriteOnlyCell(ws, value=name)
        cell.fill, cell.font = _HDR_FILL, _HDR_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        out.append(cell)
    return out


README_LINES = [
    ("SEBI Portfolio Manager Monthly Reports - extracted tables", True),
    ("", False),
    ("Source", True),
    ("https://www.sebi.gov.in/sebiweb/other/OtherAction.do?doPmr=yes", False),
    ("", False),
    ("IMPORTANT - coverage caveat", True),
    ("SEBI published a different monthly-report format before February 2021.", False),
    ("Tables B, C, G and H (the investment-approach break-ups requested here)", False),
    ("DO NOT EXIST on reports for January 2021 and earlier - those pages carry an", False),
    ("older layout with no per-approach detail. Data therefore starts 2021-02.", False),
    ("Months before that are recorded in 'Coverage' as 'Legacy Format'.", False),
    ("", False),
    ("Sheets", True),
    ("B_Disc_AUM      Table B - AUM break-up, DISCRETIONARY, per investment approach", False),
    ("C_Disc_NetFlow  Table C - Net Inflow(+)/Outflow(-) during the month only", False),
    ("G_NonDisc_AUM   Table G - AUM break-up, NON-DISCRETIONARY (one row per month)", False),
    ("H_NonDisc_Flow  Table H - Funds Inflow/Outflow DURING THE MONTH group only", False),
    ("Data_Quality    Filed rows whose components do not sum to SEBI's stated Total", False),
    ("Coverage        Per-month fetch outcome, so gaps are visible rather than silent", False),
    ("", False),
    ("Notes", True),
    ("All monetary figures are in INR crores, exactly as SEBI publishes them.", False),
    ("'Is Total Row' marks SEBI's own Total line - filter it out before summing.", False),
    ("Managers are keyed on Registration No; the newest published name is used", False),
    ("throughout (e.g. IIFL Asset Management now appears as 360 ONE Asset Management).", False),
    ("Blank numeric cells mean SEBI left the field empty, not zero.", False),
    ("Some early-2021 filings do not self-reconcile; see the Data_Quality sheet.", False),
]


def _write_readme(wb, sheets, generated_at):
    """README first, so it is the sheet that opens - written inline, not re-opened.

    The workbook can reach several hundred thousand rows; re-opening it just to
    restyle would cost gigabytes, so every sheet is styled as it streams out.
    """
    from openpyxl.cell import WriteOnlyCell
    ws = wb.create_sheet("README")
    lines = list(README_LINES)
    lines.insert(4, (f"Generated (UTC): {generated_at}", False))
    lines.append(("", False))
    lines.append(("Row counts", True))
    for name, (_, rows) in sheets.items():
        lines.append((f"{name}: {len(rows):,} rows", False))
    for text, bold in lines:
        cell = WriteOnlyCell(ws, value=text)
        if bold:
            cell.font = Font(bold=True, size=11)
        ws.append([cell])
    ws.column_dimensions["A"].width = 96


def write_workbook(conn, path: str, csv_dir: str | None = None) -> dict:
    from datetime import datetime, timezone

    sheets = collect(conn)
    head_cov, rows_cov = coverage_rows(conn)
    sheets_out = dict(sheets)
    sheets_out["Coverage"] = (head_cov, rows_cov)
    generated = datetime.now(timezone.utc).isoformat(timespec="seconds")

    if csv_dir:
        os.makedirs(csv_dir, exist_ok=True)
        for name, (head, rows) in sheets_out.items():
            with open(os.path.join(csv_dir, f"{name}.csv"), "w", newline="", encoding="utf-8") as fh:
                w = csv.writer(fh)
                w.writerow(head)
                w.writerows(rows)
        log.info("wrote %d CSVs to %s", len(sheets_out), csv_dir)

    wb = Workbook(write_only=True)
    _write_readme(wb, sheets_out, generated)
    for name, (head, rows) in sheets_out.items():
        for part, start in enumerate(range(0, max(len(rows), 1), SHEET_ROW_LIMIT)):
            title = (name if part == 0 else f"{name}_pt{part + 1}")[:31]
            ws = wb.create_sheet(title)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(max(len(head), 1))}1"
            ws.append(_header_cells(ws, head))
            for r in rows[start:start + SHEET_ROW_LIMIT]:
                ws.append(r)
            if len(rows) > SHEET_ROW_LIMIT:
                log.warning("%s split at row %d (Excel sheet limit)", name, SHEET_ROW_LIMIT)

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    log.info("wrote %s", path)
    return {name: len(rows) for name, (_, rows) in sheets_out.items()}
