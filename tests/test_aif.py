"""Tests for the AIF Category III scraper - offline, against a captured page.

The fixture is SEBI's AIF statistics page as fetched on 2026-09-28: every
quarter from September 2012 to June 2026, including the four September
sections that SEBI heads "December 31".
"""
from __future__ import annotations

import csv
import gzip
import os
import re
import sys
from datetime import date

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from sebi_aif.excel import WORKBOOK, build_tables, financial_year, fy_quarter, write_outputs
from sebi_aif.parse import (EQUITY_DEBT, NET_FIGURES, FormatError, heading_date, parse_page,
                            previous_quarter)
from sebi_pmr.tables import expand, to_number

FIXTURE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures",
                       "aif_statistics_2026_06.html.gz")
URL_FRAGMENT = "1392982252002.html"


@pytest.fixture(scope="module")
def page() -> str:
    with gzip.open(FIXTURE, "rt", encoding="utf-8") as fh:
        return fh.read()


@pytest.fixture(scope="module")
def quarters(page):
    return parse_page(page)


def section(when: str, cat3=("30", "20", "10"), extra_col: str | None = None) -> str:
    """A minimal quarter section in the page's current markup."""
    extra_h = f"<td>{extra_col}</td>" if extra_col else ""
    extra_v = "<td>7</td>" if extra_col else ""
    return (f"<strong>Cumulative net figures as at the end of the quarter ending {when}</strong>"
            "<table><tr><td colspan=4>(All figures in Rs. Crores)</td></tr>"
            "<tr><td>Category of AIF</td><td>Commitments Raised</td><td>Funds Raised</td>"
            f"<td>Investments Made</td>{extra_h}</tr>"
            f"<tr><td>Category II AIF</td><td>99</td><td>99</td><td>99</td>{extra_v}</tr>"
            f"<tr><td>Category III AIF</td><td>{cat3[0]}</td><td>{cat3[1]}</td>"
            f"<td>{cat3[2]}</td>{extra_v}</tr></table>")


# ------------------------------------------------------------- the real page
def test_every_quarter_from_2012_to_2026_without_gaps(quarters):
    ends = [q.quarter_end for q in quarters]
    assert len(ends) == 56
    assert ends[0] == date(2012, 9, 30) and ends[-1] == date(2026, 6, 30)
    assert all(previous_quarter(newer) == older for older, newer in zip(ends, ends[1:]))


def test_repeated_december_headings_are_dated_as_september(quarters):
    redated = [q for q in quarters if q.stated_date != q.quarter_end]
    assert [q.quarter_end for q in redated] == [date(y, 9, 30) for y in (2020, 2021, 2022, 2023)]
    for q in redated:
        assert q.stated_date == date(q.quarter_end.year, 12, 31)
        assert "repeats the quarter above it" in q.notes[0]


def test_redating_agrees_with_the_grand_total(page):
    """Industry-wide commitments rise in every section, newest first, so page
    order - not the repeated heading - is what dates a section.  Independently,
    Business Standard reported Rs 9.54 trn of AIF commitments at September 2023:
    the Grand Total of the second 'December 31, 2023' section."""
    html = re.sub(r"<!--.*?-->", "", page, flags=re.S)
    starts = [m.start() for m in re.finditer(r"Cumulative net figures as at the end of", html)]
    totals = []
    for i, a in enumerate(starts):
        chunk = html[a:starts[i + 1] if i + 1 < len(starts) else len(html)]
        grid = expand(re.search(r"<table\b.*?</table\s*>", chunk, re.S | re.I).group(0))
        totals.append(to_number(next(r for r in grid if r[0] == "Grand Total")[1]))
    assert all(newer > older for newer, older in zip(totals, totals[1:]))
    assert round(totals[11]) == 954397


def test_published_figures_are_reproduced_exactly(quarters):
    q = {x.quarter_end: x for x in quarters}
    assert q[date(2026, 6, 30)].net == {"Commitments Raised": 337192, "Funds Raised": 217227,
                                        "Investments Made": 230911}
    assert q[date(2023, 9, 30)].net == {"Commitments Raised": 103502.84, "Funds Raised": 71748.83,
                                        "Investments Made": 78686.59}
    assert q[date(2019, 12, 31)].net["Commitments Raised"] == 48151.361
    assert q[date(2012, 9, 30)].net == dict.fromkeys(NET_FIGURES, 0)


def test_the_category_iii_row_is_never_category_ii(quarters):
    # "Category III" begins with "Category II"; the printed labels prove which row was read.
    assert {q.net_label for q in quarters} == {"Category III", "Category III AIF"}
    assert {q.equity_debt_label for q in quarters if q.equity_debt} == {"Category III AIF"}


def test_equity_and_debt_table_exists_from_march_2024(quarters):
    with_table = [q.quarter_end for q in quarters if q.equity_debt is not None]
    assert len(with_table) == 10 and with_table[0] == date(2024, 3, 31)
    latest = quarters[-1]
    assert list(latest.equity_debt) == list(EQUITY_DEBT)
    assert list(latest.equity_debt.values()) == [183906, 2290, 40809, None]   # '-' is nil
    assert latest.footnote.startswith("* Investments made in equity")


def test_breakdown_never_exceeds_investments_made(quarters):
    for q in quarters:
        if q.equity_debt:
            assert sum(v or 0 for v in q.equity_debt.values()) <= q.net["Investments Made"]


def test_outputs_are_written(quarters, tmp_path):
    counts = write_outputs(quarters, str(tmp_path))
    assert counts == {"CatIII_NetFigures": 56, "CatIII_EquityDebt": 10, "CatIII_Combined": 56}
    with open(tmp_path / "csv" / "CatIII_Combined.csv", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert rows[-1]["Quarter End"] == "2026-06-30" and rows[-1]["Financial Year"] == "2026-27"
    assert rows[-1]["Commitments Raised (Rs crore)"] == "337192"
    assert rows[0]["Security Receipts (Rs crore)"] == ""        # before the table existed
    wb = load_workbook(tmp_path / WORKBOOK)
    assert wb.sheetnames == ["README", *counts]
    readme = "\n".join(str(c.value) for c in wb["README"]["A"] if c.value)
    assert "September 30, 2023" in readme and URL_FRAGMENT in readme



# -------------------------------------------------------- synthetic layouts
def test_commented_out_sections_are_ignored():
    page = "<!-- " + section("September 30, 2026", ("1", "1", "1")) + " -->" + section("June 30, 2026")
    (q,) = parse_page(page)
    assert q.quarter_end == date(2026, 6, 30) and q.net["Commitments Raised"] == 30


def test_a_gap_in_the_quarters_fails_loudly():
    with pytest.raises(FormatError, match="does not follow"):
        parse_page(section("June 30, 2026") + section("December 31, 2025"))


def test_a_repeat_is_only_redated_into_an_unclaimed_quarter():
    with pytest.raises(FormatError, match="does not follow"):
        parse_page(section("June 30, 2026") + section("June 30, 2026") + section("March 31, 2026"))


def test_a_column_sebi_adds_later_is_kept_without_a_unit():
    (q,) = parse_page(section("June 30, 2026", extra_col="Number of Schemes"))
    assert q.net["Number of Schemes"] == 7
    head, _ = build_tables([q])["CatIII_NetFigures"]
    assert "Number of Schemes" in head and "Funds Raised (Rs crore)" in head


def test_a_missing_column_fails_loudly():
    page = section("June 30, 2026").replace("Funds Raised", "Something Else")
    with pytest.raises(FormatError, match="Funds Raised"):
        parse_page(page)


def test_unreadable_figures_fail_and_odd_grouping_is_flagged():
    with pytest.raises(FormatError, match="unreadable"):
        parse_page(section("June 30, 2026", cat3=("12a", "1", "1")))
    (q,) = parse_page(section("June 30, 2026", cat3=("5,63", "1", "1")))
    assert q.net["Commitments Raised"] == 563
    assert "irregular digit grouping" in q.notes[0]


@pytest.mark.parametrize("text, expected", [
    ("as at the end of the quarter ending March 31 st , 2024", date(2024, 3, 31)),
    ("as at the end of 30th September 2018 (All figures in Rs. Crores)", date(2018, 9, 30)),
    ("as at the end of 31 March, 2020", date(2020, 3, 31)),
    ("the quarter ending Dec 31, 2024", date(2024, 12, 31)),
    ("as at the end of 30 June 2019", date(2019, 6, 30)),
    ("no date in this heading", None),
])
def test_heading_dates_in_every_format_seen(text, expected):
    assert heading_date(text)[0] == expected


@pytest.mark.parametrize("d, fy, fq", [
    (date(2026, 6, 30), "2026-27", "Q1"), (date(2012, 9, 30), "2012-13", "Q2"),
    (date(2012, 12, 31), "2012-13", "Q3"), (date(2026, 3, 31), "2025-26", "Q4"),
    (date(1999, 12, 31), "1999-00", "Q3"),
])
def test_indian_financial_year(d, fy, fq):
    assert (financial_year(d), fy_quarter(d)) == (fy, fq)
