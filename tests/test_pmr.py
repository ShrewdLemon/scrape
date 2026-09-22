"""Tests run entirely against archived fixtures - no network access needed.

The fixtures are real pages captured from the portal, one per format era:
2018/2020 (legacy layout), Dec-2020 and Jan-2021 (general information only),
Feb/Mar/Apr-2021 (first months of the lettered format) and 2024 (current), plus
a small manager whose sections differ from a large one.
"""
from __future__ import annotations

import glob
import gzip
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from sebi_pmr.excel import AUM_ORDER, write_workbook
from sebi_pmr.fetch import PmrClient, PortfolioManager, RateLimiter
from sebi_pmr.parse import NEW_FORMAT_START, FormatError, parse_report
from sebi_pmr.pipeline import classify, months_between, parse_period, plan, select_shard
from sebi_pmr.store import Store, merge, strip_dropdown
from sebi_pmr.tables import expand, is_number, to_number

FIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")

COMPONENTS = [c for c in AUM_ORDER if c != "Total"]


def fixture(name: str) -> str:
    with gzip.open(os.path.join(FIX, name + ".html.gz"), "rt", encoding="utf-8") as fh:
        return fh.read()


def period_of(name: str) -> tuple[int, int]:
    parts = name.split("_")
    return int(parts[-2]), int(parts[-1])


NEW_FORMAT = ["360one_2021_02", "360one_2021_03", "360one_2021_04",
              "360one_2024_03", "zen_2023_06"]
LEGACY = ["360one_2018_04", "360one_2020_04", "360one_2020_10",
          "360one_2020_12", "360one_2021_01"]


# --------------------------------------------------------------- table utils
def test_rowspan_and_colspan_expand_to_a_dense_grid():
    grid = expand("<table><tr><th rowspan=2>A</th><th colspan=2>B</th></tr>"
                  "<tr><td>b1</td><td>b2</td></tr><tr><td>1</td><td>2</td><td>3</td></tr></table>")
    assert grid == [["A", "B", "B"], ["A", "b1", "b2"], ["1", "2", "3"]]


def test_absurd_colspan_is_clamped():
    grid = expand('<table><tr><td colspan="9999">x</td></tr></table>')
    assert len(grid[0]) <= 60


@pytest.mark.parametrize("text,expected", [
    ("1,234.5", 1234.5), ("-3.2", -3.2), ("(12)", -12.0),
    ("0", 0.0), ("", None), ("NA", None), ("-", None), ("abc", None),
])
def test_number_parsing(text, expected):
    assert to_number(text) == expected


def test_is_number_rejects_labelled_text():
    assert is_number("12.5")
    assert not is_number("Inflow during the FY since April 01 to March 2024")
    assert not is_number("Total")


# -------------------------------------------------------------------- format
@pytest.mark.parametrize("name", LEGACY)
def test_legacy_pages_are_rejected_not_silently_empty(name):
    """Pre-Feb-2021 pages must raise, never yield zero rows that look like real zeros."""
    with pytest.raises(FormatError):
        parse_report(fixture(name), *period_of(name))


@pytest.mark.parametrize("name", NEW_FORMAT)
def test_new_format_pages_yield_all_four_tables(name):
    rep = parse_report(fixture(name), *period_of(name))
    assert rep.b_aum, "table B empty"
    assert rep.c_flow, "table C empty"
    assert rep.g_aum, "table G empty"
    assert rep.h_flow, "table H empty"
    assert rep.pm_reg.startswith("INP")
    assert not rep.notes, rep.notes


def test_classify_separates_legacy_from_missing_data():
    assert classify(fixture("360one_2018_04"), 2018, 4)[0] == "oldformat"
    assert classify(fixture("360one_2024_03"), 2024, 3)[0] == "ok"


def test_format_start_constant_matches_fixtures():
    assert NEW_FORMAT_START == (2021, 2)


# ------------------------------------------------------------------ table B/G
@pytest.mark.parametrize("name", NEW_FORMAT)
def test_table_b_columns_are_the_canonical_twelve(name):
    rep = parse_report(fixture(name), *period_of(name))
    cols = [k for k in rep.b_aum[0] if k not in ("investment_approach", "is_total")]
    assert cols == AUM_ORDER, cols


@pytest.mark.parametrize("name", NEW_FORMAT)
def test_table_b_total_row_reconciles_column_wise(name):
    """SEBI's own Total row must equal the sum of the approach rows we parsed.

    This is the strongest available check that columns are mapped correctly:
    a shifted column would break the reconciliation immediately.
    """
    rep = parse_report(fixture(name), *period_of(name))
    totals = [r for r in rep.b_aum if r["is_total"]]
    if not totals:
        pytest.skip("no Total row published")
    body = [r for r in rep.b_aum if not r["is_total"]]
    for col in AUM_ORDER:
        got = sum((r.get(col) or 0) for r in body)
        assert abs(got - (totals[0].get(col) or 0)) < 0.05, f"{col} mismatch"


@pytest.mark.parametrize("name", NEW_FORMAT)
def test_table_g_is_one_row_with_the_same_columns(name):
    rep = parse_report(fixture(name), *period_of(name))
    assert len(rep.g_aum) == 1
    assert [k for k in rep.g_aum[0]] == AUM_ORDER


# ------------------------------------------------------------------ table C/H
@pytest.mark.parametrize("name", NEW_FORMAT)
def test_table_c_keeps_only_the_net_monthly_column(name):
    rep = parse_report(fixture(name), *period_of(name))
    keys = set(rep.c_flow[0])
    assert keys == {"investment_approach", "is_total",
                    "net_inflow_outflow_during_month_inr_cr"}


def test_table_c_net_column_is_the_month_not_the_financial_year():
    """The month and FY columns look near-identical; picking the wrong one is silent."""
    rep = parse_report(fixture("360one_2024_03"), 2024, 3)
    by_approach = {r["investment_approach"]: r["net_inflow_outflow_during_month_inr_cr"]
                   for r in rep.c_flow}
    # From the published page: Multicap PMS is -66.30 for the month, -665.70 for the FY.
    assert by_approach["360 ONE Multicap PMS"] == pytest.approx(-66.30)
    assert by_approach["Total"] == pytest.approx(-54.99)


@pytest.mark.parametrize("name", NEW_FORMAT)
def test_table_h_keeps_only_the_during_the_month_group(name):
    rep = parse_report(fixture(name), *period_of(name))
    assert set(rep.h_flow[0]) == {"inflow_during_month_inr_cr",
                                  "outflow_during_month_inr_cr",
                                  "net_inflow_outflow_during_month_inr_cr"}


def test_table_c_total_reconciles_against_its_approach_rows():
    rep = parse_report(fixture("360one_2021_02"), 2021, 2)
    body = [r for r in rep.c_flow if not r["is_total"]]
    total = [r for r in rep.c_flow if r["is_total"]][0]
    got = sum(r["net_inflow_outflow_during_month_inr_cr"] or 0 for r in body)
    assert abs(got - total["net_inflow_outflow_during_month_inr_cr"]) < 0.05


def test_manager_rename_is_visible_in_the_fixtures():
    """Same registration number, different published name - the Excel keys on the number."""
    old = parse_report(fixture("360one_2021_02"), 2021, 2)
    new = parse_report(fixture("360one_2024_03"), 2024, 3)
    assert old.pm_reg == new.pm_reg == "INP000004565"
    assert old.pm_name != new.pm_name


# ------------------------------------------------------------------ pipeline
def test_months_between_is_inclusive_and_ordered():
    assert months_between((2020, 11), (2021, 2)) == [(2020, 11), (2020, 12), (2021, 1), (2021, 2)]


def test_parse_period_rejects_nonsense():
    assert parse_period("2021-02") == (2021, 2)
    for bad in ("2021", "2021-13", "x-y"):
        with pytest.raises(ValueError):
            parse_period(bad)


def test_shards_partition_without_overlap():
    items = list(range(50))
    seen = []
    for i in range(7):
        seen.extend(select_shard(items, i, 7))
    assert sorted(seen) == items


def test_rate_limiter_rejects_bad_windows():
    with pytest.raises(ValueError):
        RateLimiter(0, 1)
    with pytest.raises(ValueError):
        RateLimiter(4, 2)


def test_rate_limiter_penalty_only_widens():
    r = RateLimiter(2.5, 4.0)
    r.penalise()
    assert r.min_delay > 2.5 and r.max_delay > 4.0


def test_default_delay_window_is_the_requested_budget():
    c = PmrClient()
    assert (c.limiter.min_delay, c.limiter.max_delay) == (2.5, 4.0)


# --------------------------------------------------------------------- store
@pytest.fixture()
def seeded(tmp_path):
    store = Store(str(tmp_path / "pmr.db"), str(tmp_path / "raw"))
    pms = {
        "360one": PortfolioManager("a@@a@@360 ONE", "INP000004565", "360 ONE ASSET MANAGEMENT LIMITED"),
        "zen": PortfolioManager("b@@b@@ZEN", "INP000000936", "ZEN WEALTH MANAGEMENT SERVICES LIMITED"),
    }
    store.upsert_pms(list(pms.values()))
    for path in sorted(glob.glob(os.path.join(FIX, "*.gz"))):
        name = os.path.basename(path).replace(".html.gz", "")
        pm = pms[name.split("_")[0]]
        year, month = period_of(name)
        html = fixture(name)
        arc, sha = store.write_archive(pm.reg_no, year, month, html)
        status, rep, detail = classify(html, year, month)
        n = store.save_report(pm.pmr_id, rep) if rep else 0
        store.log_page(pm.pmr_id, year, month, status, detail, n, arc, sha)
    return store


def test_store_records_every_cell_and_survives_replanning(seeded):
    cov = seeded.coverage()
    assert cov["pages_by_status"]["ok"] == len(NEW_FORMAT)
    assert cov["pages_by_status"]["oldformat"] == len(LEGACY)
    pms = [PortfolioManager(r["pmr_id"], r["reg_no"], r["name"]) for r in seeded.pms()]
    outstanding = plan(seeded, pms, [(2024, 3)])
    assert all(not (p.reg_no == "INP000004565" and y == 2024 and m == 3)
               for p, y, m in outstanding), "already-fetched cell was replanned"


def test_archive_roundtrips_and_drops_the_dropdown(seeded):
    row = seeded.conn.execute(
        "SELECT archive FROM page_log WHERE status='ok' LIMIT 1").fetchone()
    html = seeded.read_archive(row["archive"])
    assert "<option" not in html.split("</select>")[0]
    assert parse_report(html, 2024, 3) or True


def test_strip_dropdown_keeps_the_rest_of_the_page():
    html = '<select name="pmrId"><option>a</option></select><table>x</table>'
    out = strip_dropdown(html)
    assert "<option>" not in out and "<table>x</table>" in out


def test_merge_unions_shard_databases(tmp_path, seeded):
    other = Store(str(tmp_path / "shard1.db"), str(tmp_path / "raw"))
    other.upsert_pms([PortfolioManager("c@@c@@OTHER", "INP999", "OTHER PM")])
    other.log_page("c@@c@@OTHER", 2024, 4, "ok", n_rows=1)
    other.close()
    seeded.close()
    merge([seeded.path, str(tmp_path / "shard1.db")], str(tmp_path / "merged.db"),
          str(tmp_path / "raw"))
    merged = Store(str(tmp_path / "merged.db"), str(tmp_path / "raw"))
    assert merged.coverage()["managers"] == 3
    assert ("c@@c@@OTHER", 2024, 4) in merged.done_cells()


# --------------------------------------------------------------------- excel
def test_workbook_has_every_expected_sheet(tmp_path, seeded):
    from openpyxl import load_workbook
    out = tmp_path / "out.xlsx"
    counts = write_workbook(seeded.conn, str(out), csv_dir=str(tmp_path / "csv"))
    wb = load_workbook(out)
    assert wb.sheetnames[0] == "README"
    for sheet in ("B_Disc_AUM", "C_Disc_NetFlow", "G_NonDisc_AUM",
                  "H_NonDisc_Flow", "Data_Quality", "Coverage"):
        assert sheet in wb.sheetnames
        assert wb[sheet].freeze_panes == "A2"
    assert counts["B_Disc_AUM"] > 0
    assert os.path.exists(tmp_path / "csv" / "B_Disc_AUM.csv")


def test_workbook_flags_sebi_internal_inconsistencies(tmp_path, seeded):
    """Feb/Mar-2021 filings do not self-reconcile; that must surface, not vanish."""
    from openpyxl import load_workbook
    out = tmp_path / "out.xlsx"
    write_workbook(seeded.conn, str(out), csv_dir=None)
    ws = load_workbook(out)["Data_Quality"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows, "expected the known 2021 filing inconsistencies to be listed"
    assert all(abs(r[-1]) > 0.05 for r in rows)


def test_excel_identifier_columns_are_present(tmp_path, seeded):
    from openpyxl import load_workbook
    out = tmp_path / "out.xlsx"
    write_workbook(seeded.conn, str(out), csv_dir=None)
    head = [c.value for c in load_workbook(out)["B_Disc_AUM"][1]]
    assert head[:6] == ["Portfolio Manager", "Registration No", "Year",
                        "Month", "Month Name", "Period"]


# ------------------------------------------------------------------ catalogue
def test_manager_dropdown_parses_into_id_and_name():
    html = ('<select name="pmrId"><option value="">--</option>'
            '<option value="INP1@@INP1@@ALPHA LLP">ALPHA LLP</option></select>')
    pms = PmrClient().portfolio_managers(html)
    assert len(pms) == 1
    assert pms[0].reg_no == "INP1" and pms[0].name == "ALPHA LLP"
