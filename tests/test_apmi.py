"""APMI IA Insights matching - offline, on names seen in the real data."""
from __future__ import annotations

import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from sebi_pmr import apmi


def ia(ia_id, name, manager):
    return {"id": ia_id, "iaName": name, "pmsName": manager}


def test_parse_rsc_returns_value_and_raises_on_error():
    ok = '0:{"a":"$@1","f":"","b":"x"}\n1:[{"iaId":"a"}]\n'
    assert apmi.parse_rsc(ok) == [{"iaId": "a"}]
    with pytest.raises(apmi.NoData):
        apmi.parse_rsc('0:{"a":"$@1"}\n1:E{"digest":"4155642546"}\n')
    with pytest.raises(apmi.NoData):
        apmi.parse_rsc("<!DOCTYPE html><html>")


@pytest.mark.parametrize("apmi_name,sebi_name", [
    ("360 One Portfolio Managers Limited (Formerly Known As Iifl Wealth Portfolio "
     "Managers Limited)", "IIFL WEALTH PORTFOLIO MANAGERS LIMITED"),
    ("360 One Portfolio Managers Limited (Formerly Known As Iifl Wealth Portfolio "
     "Managers Limited)", "360 ONE PORTFOLIO MANAGERS LIMITED"),
    ("Molecule Ventures Limited Liability Partnership", "Molecule Ventures LLP"),
    ("Motilal Oswal Asset Management Company Limited - Portfolio Managers",
     "MOTILAL OSWAL ASSET MANAGEMENT COMPANY LIMITED"),
    ("Tamohara Investment Managers Pvt Ltd", "TAMOHARA INVESTMENT MANAGERS PRIVATE LIMITED"),
])
def test_manager_names_map_to_registrations(apmi_name, sebi_name):
    regs = apmi.provider_regs({apmi_name}, {"INP1": [sebi_name]})
    assert regs[apmi_name] == {"INP1"}


def test_manager_with_two_registrations_maps_to_both():
    regs = apmi.provider_regs({"ASK Wealth Advisors Private Limited"},
                              {"INP7": ["ASK WEALTH ADVISORS PRIVATE LIMITED"],
                               "INP8": ["ASK Wealth Advisors Private Limited"]})
    assert regs["ASK Wealth Advisors Private Limited"] == {"INP7", "INP8"}


@pytest.mark.parametrize("a,b", [
    ("Motilal Oswal Multicap Oppurtunities Strategy.", "Motilal Oswal Multicap Opportunities Strategy"),
    ("SHREE DYNAMOIC PLAN", "Shree Dynamic Plan"),
    ("EMKAYS GOLDEN", "EMKAYGOLDEN"),
    ("RH ALLIANCE PORTFOLIO", "RHPMPL - ALLIANCE PORTFOLIO"),
    ("Large Cap Blue Chip Fund", "Val-Q Large Cap Blue Chip Fund"),
    ("Mehta Multi Focus Strategy", "MEHTA MULTI FOCUS STRATEGY (MMFS)"),
])
def test_fuzzy_accepts_typos_and_prefixes(a, b):
    managers = {"Val-Q Large Cap Blue Chip Fund": "Val Q Investment Advisory Pvt Ltd",
                "RHPMPL - ALLIANCE PORTFOLIO": "Right Horizons Portfolio Management Pvt Ltd"}
    assert apmi.fuzzy_score(a, b, managers.get(b)) > 0


@pytest.mark.parametrize("a,b", [
    ("Barclays Discretionary Short Duration", "Barclays Non-Discretionary Short Duration"),
    ("KINETIC 2", "KINETIC"),
    ("Q India Value Equity Strategy - Constrained XIII", "Q India Value Equity Strategy - Constrained XVIII"),
    ("RENAISSANCE MULTICAP PORTFOLIO", "Renaissance Midcap Portfolio"),
    ("Approach 1 Equity Opportunity", "Approach 1 Elite Opportunity"),
    ("R Wadiwala Shariah Agile Pe Ratio Portfolio", "R Wadiwala Agile PE Ratio Portfolio"),
    ("ICICI Prudential PMS Multi-Manager Strategy", "ICICI Prudential PMS Multi-Manager - Thematic Strategy"),
])
def test_fuzzy_rejects_different_products(a, b):
    assert apmi.fuzzy_score(a, b) == 0


def _match(unique, regs_by_name, universe, managers, ranked=frozenset()):
    prov = apmi.provider_regs({a["pmsName"] for a in universe}, managers)
    return apmi.match_all(unique, regs_by_name, {r: n[0] for r, n in managers.items()},
                          universe, prov, ranked=set(ranked))


def test_same_name_different_managers_each_get_their_own_approach():
    universe = [ia("a", "Money Multiplier", "Alpha Capital Pvt Ltd"),
                ia("b", "Money Multiplier", "Beta Advisors LLP")]
    managers = {"INP1": ["ALPHA CAPITAL PRIVATE LIMITED"], "INP2": ["Beta Advisors LLP"]}
    rows = _match(["Money Multiplier"], {"Money Multiplier": {"INP1", "INP2"}}, universe, managers)
    assert [(m.reg_no, m.apmi["id"], m.how) for m in rows] == [
        ("INP1", "a", "Exact (name + manager)"), ("INP2", "b", "Exact (name + manager)")]


def test_name_from_another_manager_is_not_matched():
    universe = [ia("a", "Money Multiplier", "Alpha Capital Pvt Ltd")]
    managers = {"INP1": ["Alpha Capital Pvt Ltd"], "INP2": ["Beta Advisors LLP"]}
    rows = _match(["Money Multiplier"], {"Money Multiplier": {"INP2"}}, universe, managers)
    assert rows[0].how == "Not on APMI" and rows[0].apmi is None


def test_discretionary_name_prefers_non_nd_candidate():
    universe = [ia("d", "Buoyant Opportunities PMS", "Buoyant Capital Private Limited"),
                ia("n", "Buoyant Opportunities NDPMS", "Buoyant Capital Private Limited")]
    managers = {"INP1": ["Buoyant Capital Private Limited"]}
    rows = _match(["Buoyant Opportunities Scheme"], {"Buoyant Opportunities Scheme": {"INP1"}},
                  universe, managers)
    assert [m.apmi["id"] for m in rows] == ["d"]
    assert rows[0].how == "Normalised (name + manager)"


def test_apmi_duplicate_resolved_by_ranking_and_spelling():
    universe = [ia("x", "Vireya - Growth Anchors", "Vireya Capital Management LLP"),
                ia("y", "Vireya - Growth Anchors", "Vireya Capital Management LLP"),
                ia("s1", "Multi-Asset Allocation", "Buglerock Capital Private Limited"),
                ia("s2", "Multi Asset Allocation", "Buglerock Capital Private Limited")]
    managers = {"INP1": ["Vireya Capital Management LLP"], "INP2": ["Buglerock Capital Pvt Ltd"]}
    rows = _match(["Vireya - Growth Anchors", "Multi Asset Allocation"],
                  {"Vireya - Growth Anchors": {"INP1"}, "Multi Asset Allocation": {"INP2"}},
                  universe, managers, ranked={"y"})
    assert [m.apmi["id"] for m in rows] == ["y", "s2"]


def test_nd_names_match_on_name_alone_and_placeholders_are_marked():
    universe = [ia("a", "ASK Alchemist", "Ask Wealth Advisors Private Limited")]
    rows = _match(["ASK Alchemist", "0", "NA"], {}, universe, {})
    assert [m.how for m in rows] == ["Exact (name only)", "SEBI placeholder", "SEBI placeholder"]


def test_fuzzy_row_names_its_exact_sibling():
    universe = [ia("a", "Shree Dynamic Plan", "Shree Rama Managers Llp")]
    managers = {"INP1": ["Shree Rama Managers LLP"]}
    rows = _match(["Shree Dynamic Plan", "SHREE DYNAMOIC PLAN"],
                  {"Shree Dynamic Plan": {"INP1"}, "SHREE DYNAMOIC PLAN": {"INP1"}},
                  universe, managers)
    assert rows[1].how.startswith("Fuzzy") and rows[1].apmi["id"] == "a"
    assert '"Shree Dynamic Plan"' in rows[1].note


def test_write_sheet_fills_details_and_links():
    from openpyxl import Workbook
    wb = Workbook()
    rows = [apmi.Match("Alpha", "INP1", "Alpha Capital", "Exact (name + manager)",
                       ia("a", "Alpha", "Alpha Capital Pvt Ltd")),
            apmi.Match("Beta", "INP2", "Beta LLP", "Not on APMI")]
    details = {"a": {"strategy": "Equity", "serviceType": "N", "dateOfInception": "13-08-2025"}}
    ws = apmi.write_sheet(wb, rows, details, "23-Sep-2026")
    assert [c.value for c in ws[5]][:6] == ["Alpha", "Alpha Capital", "INP1", "Equity",
                                            "Non-Discretionary", apmi.inception("13-08-2025")]
    assert ws.cell(5, 10).hyperlink.target == apmi.REPORT_URL.format("a")
    assert ws.cell(6, 4).value is None and ws.cell(6, 9).value == "Not on APMI"
    assert ws.auto_filter.ref == "A4:K6"
