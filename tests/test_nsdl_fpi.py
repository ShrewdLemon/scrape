"""NSDL FPI equity scraper - offline, against real captured report pages.

Fixtures cover both table eras: 2002/2013 (Equity, Debt, Total) and 2026
(eleven columns, including a second "Equity" under Mutual Funds that must not
be confused with the FPI Equity column), in INR and USD.
"""
from __future__ import annotations

import gzip
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from nsdl_fpi.excel import long_rows, mismatches, write_workbook
from nsdl_fpi.fetch import available_years, hidden_fields
from nsdl_fpi.parse import FormatError, parse_year, to_number

FIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "nsdl")


def page(name: str) -> str:
    with gzip.open(os.path.join(FIX, name + ".html.gz"), "rt", encoding="utf-8") as fh:
        return fh.read()


ALL = [("2002_INR", 2002, "INR"), ("2002_USD", 2002, "USD"), ("2013_INR", 2013, "INR"),
       ("2026_INR", 2026, "INR"), ("2026_USD", 2026, "USD")]


@pytest.mark.parametrize("name,year,cur", ALL)
def test_months_reconcile_with_nsdl_total(name, year, cur):
    rep = parse_year(page(name), year, cur)
    assert rep.stated_total is not None
    assert rep.computed_total == rep.stated_total


def test_old_layout_values():
    rep = parse_year(page("2002_INR"), 2002, "INR")
    assert rep.leaf_headers == ["Equity", "Debt", "Total"]
    assert rep.equity == {1: 424, 2: 1967, 3: 391, 4: 10, 5: -55, 6: -382,
                          7: 350, 8: 207, 9: 468, 10: -776, 11: 601, 12: 429}


def test_new_layout_takes_fpi_equity_not_mutual_fund_equity():
    rep = parse_year(page("2026_INR"), 2026, "INR")
    assert rep.leaf_headers.count("Equity") == 2
    assert rep.equity[1] == -35962           # FPI equity; MF equity that month is 312
    assert rep.equity[3] == -117775


def test_partial_year_keeps_only_published_months():
    rep = parse_year(page("2026_INR"), 2026, "INR")
    assert sorted(rep.equity) == list(range(1, 10))


def test_usd_values():
    assert parse_year(page("2002_USD"), 2002, "USD").equity[1] == 87
    assert parse_year(page("2026_USD"), 2026, "USD").equity[1] == -3976


def test_wrong_year_or_currency_is_refused():
    with pytest.raises(FormatError):
        parse_year(page("2002_INR"), 2003, "INR")
    with pytest.raises(FormatError):
        parse_year(page("2002_INR"), 2002, "USD")


def test_missing_table_is_refused():
    with pytest.raises(FormatError):
        parse_year("<html><body>maintenance</body></html>", 2002, "INR")


def test_form_helpers():
    html = page("2026_INR")
    assert available_years(html)[0] == 2002
    assert {"__VIEWSTATE", "__EVENTVALIDATION"} <= set(hidden_fields(html))


@pytest.mark.parametrize("text,value", [("1,234", 1234), ("-55", -55), ("(12)", -12),
                                        ("3.5", 3.5), ("", None), ("&nbsp;", None)])
def test_to_number(text, value):
    assert to_number(text.replace("&nbsp;", "")) == value


def test_workbook(tmp_path):
    reps = [parse_year(page(n), y, c) for n, y, c in ALL]
    assert mismatches(reps) == []
    rows = long_rows(reps, ["INR", "USD"])
    assert len(rows) == 12 + 12 + 9          # 2002, 2013, 2026 (Jan-Sep)
    assert rows[0]["Equity (USD Million)"] == 87
    assert rows[12]["Equity (USD Million)"] is None   # 2013 USD not in fixtures

    from openpyxl import load_workbook
    out = tmp_path / "fpi.xlsx"
    write_workbook(str(out), reps, ["INR", "USD"])
    wb = load_workbook(out)
    assert wb.sheetnames == ["README", "Equity_Monthly", "INR_by_Year", "USD_by_Year", "Check"]
    assert wb["INR_by_Year"]["N2"].value == 3634
