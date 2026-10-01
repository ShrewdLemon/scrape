"""NSDL daily (Archive) equity scraper - offline, against real captured pages.

Fixtures span both layouts: 2001-01 and 2009-11 (flat Equity/Debt rows),
2009-12 (first month with Stock Exchange / Primary / Sub-total routes) and
2020-01 (routes plus Hybrid and Debt-VRR). The Total-row fallback never fires
on real 2001-2020 data, so it is exercised with small synthetic tables.
"""
from __future__ import annotations

import gzip
import os
import sys
from datetime import date

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from nsdl_fpi.daily.excel import (ROUTE_COLS, check, check_routes, daily_rows, mismatches,
                                  route_rows, write_csv, write_workbook)
from nsdl_fpi.daily.fetch import month_end, months
from nsdl_fpi.daily.parse import (BASIS_EQUITY, BASIS_ROUTES, BASIS_SUBTOTAL, BASIS_TOTAL,
                                  BASIS_TOTAL_SUM, ROUTE_NAMES, FormatError, parse_month)

FIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "nsdl_daily")
ALL = [(2001, 1), (2009, 11), (2009, 12), (2020, 1)]


def page(y: int, m: int) -> str:
    with gzip.open(os.path.join(FIX, f"{y}-{m:02d}.html.gz"), "rt", encoding="utf-8") as fh:
        return fh.read()


@pytest.mark.parametrize("y,m", ALL)
def test_daily_equity_reconciles_with_month_total(y, m):
    rep = parse_month(page(y, m), y, m)
    di, du, ok = check(rep)
    assert ok, (di, du)
    assert not rep.fallback_days
    assert all(d.day.year == y and d.day.month == m for d in rep.days)


def test_flat_layout_takes_the_equity_row():
    rep = parse_month(page(2001, 1), 2001, 1)
    assert rep.layout == "flat" and len(rep.days) == 22
    d = rep.days[0]
    assert (d.day, d.basis) == (date(2001, 1, 1), BASIS_EQUITY)
    assert (d.gross_purchases, d.gross_sales, d.net_inr, d.net_usd, d.fx) == (522.4, 189.5, 332.9, 71.2, 46.75)
    assert rep.stated[2] == 4045 and rep.computed_net_inr == 4045


def test_routed_layout_takes_the_equity_subtotal_not_stock_exchange():
    rep = parse_month(page(2009, 12), 2009, 12)
    assert rep.layout == "routed"
    d = rep.days[0]
    assert d.basis == BASIS_SUBTOTAL
    assert (d.gross_purchases, d.net_inr, d.net_usd) == (2976.1, 699.8, 150.57)
    assert {x.basis for x in rep.days} == {BASIS_SUBTOTAL}


def test_routes_captured_per_day():
    rep = parse_month(page(2009, 12), 2009, 12)
    d = rep.days[0]
    assert list(d.routes) == list(ROUTE_NAMES)
    assert d.routes["Stock Exchange"] == (2812.3, 2274.7, 537.5, 115.64)
    assert d.routes["Primary market & others"] == (163.8, 1.5, 162.3, 34.92)
    assert d.routes["Sub-total"] == (d.gross_purchases, d.gross_sales, d.net_inr, d.net_usd)
    for day in rep.days:               # Stock Exchange + Primary = Sub-total, to rounding
        se, pm, st = (day.routes[n] for n in ROUTE_NAMES)
        assert all(abs(se[i] + pm[i] - st[i]) <= 0.11 for i in range(4))


@pytest.mark.parametrize("y,m", [(2009, 12), (2020, 1)])
def test_routes_reconcile_with_month_route_totals(y, m):
    rep = parse_month(page(y, m), y, m)
    assert set(rep.stated_routes) == set(ROUTE_NAMES)
    worst, ok = check_routes(rep)
    assert ok and worst <= 0.5


def test_flat_layout_has_no_routes():
    rep = parse_month(page(2009, 11), 2009, 11)
    assert all(d.routes == {} for d in rep.days) and rep.stated_routes == {}
    assert check_routes(rep) == (None, True)
    assert route_rows([rep]) == []


def test_negatives_and_extra_categories():
    rep = parse_month(page(2020, 1), 2020, 1)
    assert rep.categories == ["Debt", "Debt-VRR", "Equity", "Hybrid"]
    d = rep.days[0]
    assert d.net_inr == -1972.18 and d.net_usd == -276.7 and d.fx == 71.274
    assert rep.computed_net_inr == rep.stated[2] == 12122.58


def test_refuses_page_for_another_month():
    with pytest.raises(FormatError):
        parse_month(page(2009, 12), 2009, 11)


def test_refuses_page_without_table():
    with pytest.raises(FormatError):
        parse_month("<html><body>maintenance</body></html>", 2010, 1)


# --- fallback to the Total row (synthetic) ----------------------------------

HEAD = ("<table class='tbls01'><tr><th colspan='8'>Daily Trends in FPI Investments up to 31-Jan-2015</th></tr>"
        "<tr><th>Reporting Date</th><th>Debt/Equity</th><th>Investment Route</th><th>Gross Purchases(Rs Crore)</th>"
        "<th>Gross Sales(Rs Crore)</th><th>Net Investment (Rs Crore)</th><th>Net Investment US($) million</th>"
        "<th>Conversion</th></tr>")


def _day(d, rows, fx="Rs.62.0000"):
    first, *rest = rows
    html = f"<tr><td rowspan='{len(rows)}'>{d}</td>{first}<td rowspan='{len(rows)}'> {fx}</td></tr>"
    return html + "".join(f"<tr>{r}</tr>" for r in rest)


def test_missing_equity_falls_back_to_total_row():
    html = HEAD + _day("02-Jan-2015", [
        "<td rowspan='3'>Equity</td><td>Stock Exchange</td><td>-</td><td>-</td><td>-</td><td>-</td>",
        "<td>Primary market & others</td><td>-</td><td>-</td><td>-</td><td>-</td>",
        "<td>Sub-total</td><td>-</td><td>-</td><td>-</td><td>-</td>",
        "<td rowspan='2'>Debt</td><td>Sub-total</td><td>10.00</td><td>5.00</td><td>5.00</td><td>0.08</td>",
        "<td>Total</td><td>110.00</td><td>15.00</td><td>95.00</td><td>1.53</td>",
    ]) + "</table>"
    d, = parse_month(html, 2015, 1).days
    assert (d.basis, d.net_inr, d.net_usd, d.fx) == (BASIS_TOTAL, 95.0, 1.53, 62.0)


def test_equity_routes_summed_when_no_subtotal():
    html = HEAD + _day("02-Jan-2015", [
        "<td rowspan='2'>Equity</td><td>Stock Exchange</td><td>100.00</td><td>40.00</td><td>60.00</td><td>0.97</td>",
        "<td>Primary market & others</td><td>1.50</td><td>0.50</td><td>1.00</td><td>0.02</td>",
        "<td>Total</td><td>101.50</td><td>40.50</td><td>61.00</td><td>0.99</td>",
    ]) + "</table>"
    d, = parse_month(html, 2015, 1).days
    assert (d.basis, d.gross_purchases, d.net_inr) == (BASIS_ROUTES, 101.5, 61.0)


def test_flat_layout_without_equity_sums_categories():
    html = HEAD + _day("02-Jan-2015", [
        "<td>Debt</td><td>10.00</td><td>4.00</td><td>6.00</td><td>0.10</td>",
        "<td>Hybrid</td><td>1.00</td><td>0.00</td><td>1.00</td><td>0.02</td>",
    ]) + "</table>"
    d, = parse_month(html, 2015, 1).days
    assert (d.basis, d.net_inr, d.net_usd) == (BASIS_TOTAL_SUM, 7.0, 0.12)


# --- helpers & output ---------------------------------------------------------

def test_month_helpers():
    assert month_end(2004, 2) == date(2004, 2, 29)
    span = months((2001, 1), (2020, 1))
    assert len(span) == 229 and span[0] == (2001, 1) and span[-1] == (2020, 1)


def test_workbook_and_csv(tmp_path):
    from openpyxl import load_workbook
    reps = [parse_month(page(y, m), y, m) for y, m in ALL]
    assert mismatches(reps) == []
    out = tmp_path / "daily.xlsx"
    write_workbook(str(out), reps)
    wb = load_workbook(out)
    assert wb.sheetnames == ["README", "Equity_Daily", "Monthly", "INR_by_Year",
                             "Equity_Routes", "Routes_Monthly", "Fallbacks"]
    rows = daily_rows(reps)
    assert wb["Equity_Daily"].max_row == len(rows) + 1 == sum(len(r.days) for r in reps) + 1
    assert wb["Equity_Daily"]["D2"].value == BASIS_EQUITY
    assert [c.value for c in wb["Monthly"][2]][-1] == "yes"
    csv_path = tmp_path / "daily.csv"
    write_csv(str(csv_path), rows)
    assert csv_path.read_text().splitlines()[1].startswith("2001-01-01,2001,1,Equity,522.4")
    rr = route_rows(reps)                # only the routed months: 2009-12 and 2020-01
    assert len(rr) == len(reps[2].days) + len(reps[3].days)
    ws = wb["Equity_Routes"]
    assert ws.max_row == len(rr) + 2     # two header rows
    assert [c.value for c in ws[3]][3:7] == [2812.3, 2274.7, 537.5, 115.64]
    assert [c.value for c in wb["Routes_Monthly"][2]][-1] == "yes"
    write_csv(str(tmp_path / "routes.csv"), rr, ROUTE_COLS)
    assert (tmp_path / "routes.csv").read_text().splitlines()[1].startswith("2009-12-01,2009,12,2812.3,2274.7,537.5")


# --- four-sheet route workbook -------------------------------------------------

def test_fy_label_and_formula():
    from nsdl_fpi.daily.route_workbook import fy_formula, fy_label
    assert fy_label(date(2010, 3, 31)) == "FY 2009-10"
    assert fy_label(date(2010, 4, 1)) == "FY 2010-11"
    assert fy_label(date(2001, 1, 1)) == "FY 2000-01"
    assert fy_formula("A5").startswith("=IF(MONTH(A5)>=4,")


def test_route_workbook_layout(tmp_path):
    from openpyxl import load_workbook
    from nsdl_fpi.daily.route_workbook import FIRST, HEADER_ROW, write_route_workbook
    reps = [parse_month(page(y, m), y, m) for y, m in ALL]
    out = tmp_path / "routes.xlsx"
    info = write_route_workbook(str(out), reps)
    n_total = sum(len(r.days) for r in reps)
    n_route = sum(1 for r in reps for d in r.days if d.routes)
    assert info["total_rows"] == n_total
    assert info["route_rows"] == {"Stock Exchange": n_route, "Primary": n_route}
    wb = load_workbook(out)
    assert wb.sheetnames == ["Total", "Stock Exchange", "Primary", "Combined", "Combined_Monthly"]
    cm = wb["Combined_Monthly"]
    assert info["months"] == len(ALL) and cm.max_row == FIRST + len(ALL)
    assert [c.value for c in cm[HEADER_ROW]][:6] == ["Month", "Year", "Month No.", "FY", "Reporting days",
                                                     "Days with route split"]
    assert cm[f"I{FIRST}"].value == f"=SUMIFS(Combined!$H$5:$H${FIRST + n_total - 1},Combined!$B$5:$B${FIRST + n_total - 1},$B5,Combined!$C$5:$C${FIRST + n_total - 1},$C5)"
    assert cm[f"A{FIRST + len(ALL)}"].value == "Total" and cm["V2"].value == 1.0
    for name in ("Total", "Stock Exchange", "Primary"):
        ws = wb[name]
        assert [c.value for c in ws[HEADER_ROW]][3] == "FY"
        assert ws[f"D{FIRST}"].value.startswith("=IF(MONTH(A5)")
    se = wb["Stock Exchange"]
    assert [se.cell(row=FIRST, column=c).value for c in range(5, 10)] == \
        ["Stock Exchange", 2812.3, 2274.7, 537.5, 115.64]
    assert wb["Primary"].cell(row=FIRST, column=8).value == 162.3
    assert wb["Total"].max_row == FIRST + n_total - 1
    cb = wb["Combined"]
    assert cb.max_row >= FIRST + n_total - 1
    assert [c.value for c in cb[HEADER_ROW]][:5] == ["Date", "Year", "Month", "FY", "Route split published?"]
    assert cb[f"A{FIRST}"].value == f"=Total!A{FIRST}"
    assert "MATCH($A5,'Stock Exchange'!$A$5:" in cb[f"V{FIRST}"].value
    assert cb[f"L{FIRST}"].value.startswith("=IF($V5=\"\",\"\",INDEX('Stock Exchange'!H$5:")
    assert cb[f"T{FIRST}"].value.count("$T$2") == 1 and cb["T2"].value == 0.1
    assert cb.column_dimensions["V"].hidden and cb.column_dimensions["W"].hidden
    heads = [c.value for c in cb[3] if c.value]
    assert "SUMMARY BY FINANCIAL YEAR (Apr-Mar)" in heads and "SUMMARY BY CALENDAR YEAR" in heads
